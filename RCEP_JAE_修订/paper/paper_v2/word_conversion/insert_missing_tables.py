from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
PATH = ROOT.parent / "main_english_v2_editable.docx"


def add_table_before(document, marker, caption, rows, category_rows=(), widths=None):
    target = next(
        p for p in document.paragraphs if p.text.strip().startswith(marker)
    )
    caption_paragraph = document.add_paragraph(caption, style="Table Caption")
    target._p.addprevious(caption_paragraph._p)

    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cells[col_index].text = str(value)
            for paragraph in cells[col_index].paragraphs:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if row_index == 0 or row_index in category_rows:
                        run.bold = True
                    if row_index in category_rows:
                        run.italic = True
        if row_index in category_rows:
            merged = cells[0]
            for cell in cells[1:]:
                merged = merged.merge(cell)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    target._p.addprevious(table._tbl)


def main():
    document = Document(PATH)

    add_table_before(
        document,
        "Notes: Statistics",
        "Summary statistics",
        [
            ["Variable", "Observations", "Mean", "S.D.", "Min.", "Median", "Max."],
            ["Relationship-year variables", "", "", "", "", "", ""],
            ["Active relationship", "630,280", "0.226", "0.418", "--", "--", "--"],
            ["RCEP partner", "630,280", "0.183", "0.386", "--", "--", "--"],
            ["Post-2022 indicator", "630,280", "0.375", "0.484", "--", "--", "--"],
            ["RCEP x Post-2022", "630,280", "0.068", "0.253", "--", "--", "--"],
            ["Foreign supplier relationship", "630,280", "0.413", "0.492", "--", "--", "--"],
            ["Relationship age in 2021 (years)", "630,280", "0.478", "1.387", "0.000", "0.000", "18.746"],
            ["Pre-policy RCEP sourcing breadth", "630,280", "0.512", "1.302", "0.000", "0.000", "8.000"],
            ["Large-firm indicator", "362,440", "0.501", "0.500", "--", "--", "--"],
            ["Partner-country-year indicators", "", "", "", "", "", ""],
            ["Incremental tariff relief (percentage points)", "1,368", "0.157", "2.490", "0.000", "0.000", "65.751"],
            ["Intermediate-goods share of bilateral imports (percent)", "1,355", "66.811", "28.800", "0.000", "73.156", "100.000"],
            ["Partner-country share of China's imports (percent)", "1,376", "0.579", "1.469", "0.000", "0.035", "10.882"],
            ["Firm-year covariates", "", "", "", "", "", ""],
            ["Log total assets", "35,978", "22.213", "1.426", "19.142", "22.021", "26.588"],
            ["Leverage", "35,978", "0.411", "0.206", "0.056", "0.399", "0.924"],
            ["Return on assets", "35,978", "0.032", "0.074", "-0.309", "0.036", "0.213"],
            ["Revenue growth", "30,283", "0.103", "0.329", "-0.613", "0.067", "1.731"],
            ["Inventory intensity", "35,337", "0.128", "0.111", "0.000", "0.104", "0.609"],
        ],
        category_rows=(1, 10, 14),
        widths=(2.7, 0.8, 0.65, 0.65, 0.65, 0.7, 0.65),
    )

    add_table_before(
        document,
        "Notes: The outcome is an active",
        "Baseline results",
        [
            ["", "(1) Common 2022 date", "(2) Member-specific dates"],
            ["", "Pair-clustered S.E.", "Pair-clustered S.E."],
            ["RCEP treatment", "0.0136***", "0.0144***"],
            ["", "(0.0040)", "(0.0041)"],
            ["Observations", "630,280", "630,280"],
            ["Relationship pairs", "78,785", "78,785"],
            ["Pair fixed effects", "Yes", "Yes"],
            ["Year fixed effects", "Yes", "Yes"],
        ],
        widths=(2.8, 1.7, 1.7),
    )

    add_table_before(
        document,
        "Notes: Each row uses its mechanism",
        "Policy-linked mechanism evidence",
        [
            ["Mechanism measure", "a", "b", "c", "c'", "Indirect effect (ab)", "Bootstrap 95% CI"],
            ["Incremental tariff relief", "0.2186***", "0.0688***", "0.0133***", "-0.0017", "0.0150***", "[0.0087, 0.0209]"],
            ["", "(0.0021)", "(0.0148)", "(0.0040)", "(0.0050)", "", ""],
            ["Intermediate-goods import share", "0.0176***", "0.0181***", "0.0151***", "0.0148***", "0.0003***", "[0.0002, 0.0004]"],
            ["", "(0.0006)", "(0.0035)", "(0.0041)", "(0.0041)", "", ""],
            ["Partner-country import share", "-0.0127***", "-2.1607***", "0.0133***", "-0.0141***", "0.0274***", "[0.0240, 0.0307]"],
            ["", "(0.0001)", "(0.1394)", "(0.0040)", "(0.0043)", "", ""],
        ],
        widths=(2.2, 0.6, 0.6, 0.6, 0.6, 1.1, 1.2),
    )

    add_table_before(
        document,
        "Notes: Group coefficients",
        "Heterogeneity analysis",
        [
            ["Partition", "Group 1", "Group 2", "Difference (Group 2 - Group 1)"],
            ["Relationship direction", "Customer: -0.0033", "Supplier: 0.0295***", "0.0327***"],
            ["", "(0.0052)", "(0.0063)", "(0.0082)"],
            ["Pre-policy relationship maturity", "Other: 0.0183***", "Long-standing: -0.0162", "-0.0345**"],
            ["", "(0.0041)", "(0.0160)", "(0.0165)"],
            ["Pre-RCEP FTA coverage", "Existing FTA: -0.0026", "Japan: 0.0378***", "0.0403***"],
            ["", "(0.0050)", "(0.0061)", "(0.0075)"],
            ["Firm sector, supplier links", "Non-manufacturing: -0.0205", "Manufacturing: 0.0274**", "0.0479**"],
            ["", "(0.0183)", "(0.0127)", "(0.0223)"],
        ],
        widths=(2.3, 1.3, 1.5, 1.4),
    )

    document.save(PATH)
    print(PATH)


if __name__ == "__main__":
    main()
