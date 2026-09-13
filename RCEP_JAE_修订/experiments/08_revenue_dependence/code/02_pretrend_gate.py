#!/usr/bin/env python3
"""Run the frozen quantified-dependence moderation pretrend gate."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from pyhdfe import create
from scipy import stats


EXPERIMENT = Path(__file__).resolve().parents[1]
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
FIGURES = EXPERIMENT / "figures"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"


plt.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": ["Times New Roman", "DejaVu Serif"],
        "font.size": 9,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 300,
    }
)


def term_name(prefix: str, group: int, event_time: int) -> str:
    suffix = f"m{abs(event_time)}" if event_time < 0 else f"p{event_time}"
    return f"{prefix}_g{group}_e_{suffix}"


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, float):
                cell.text = "" if np.isnan(value) else f"{value:.6f}"
            else:
                cell.text = str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    panel = pd.read_parquet(DATA / "quantified_relationship_year.parquet")
    if panel.duplicated(["relationship_id", "year"]).any():
        raise ValueError("Quantified relationship-year key is not unique.")
    support = json.loads((DATA / "support_gate.json").read_text(encoding="utf-8"))
    if not support["gate_pass"]:
        raise RuntimeError("The frozen data support gate did not pass.")

    work = panel.copy()
    cohort_sizes = (
        work.loc[work["rcep"].eq(1), ["relationship_id", "entry_policy_year"]]
        .drop_duplicates()
        .groupby("entry_policy_year")["relationship_id"]
        .nunique()
        .to_dict()
    )
    main_terms: list[str] = []
    interaction_terms: list[str] = []
    meta: list[dict[str, int | str]] = []
    for group in sorted(cohort_sizes):
        event_times = sorted(
            work.loc[work["entry_policy_year"].eq(group), "event_time"]
            .dropna()
            .astype(int)
            .unique()
        )
        for event_time in event_times:
            if event_time == -1:
                continue
            main_term = term_name("main", int(group), int(event_time))
            interaction_term = term_name("revenue10", int(group), int(event_time))
            indicator = (
                work["entry_policy_year"].eq(group) & work["event_time"].eq(event_time)
            ).astype("int8")
            work[main_term] = indicator
            work[interaction_term] = indicator * work["revenue_percent10"]
            main_terms.append(main_term)
            interaction_terms.append(interaction_term)
            meta.append(
                {
                    "main_term": main_term,
                    "interaction_term": interaction_term,
                    "group": int(group),
                    "event_time": int(event_time),
                }
            )

    terms = main_terms + interaction_terms
    y = work["active"].to_numpy(dtype=float)
    x = work[terms].to_numpy(dtype=float)
    fixed_effects = work[["relationship_id", "firm_year_id"]].to_numpy(dtype=np.int64)
    algorithm = create(fixed_effects, drop_singletons=False)
    residualized = algorithm.residualize(np.column_stack([y, x]))
    y_tilde, x_tilde = residualized[:, 0], residualized[:, 1:]
    clusters = pd.Categorical(work["partner_country"]).codes
    cluster_count = int(work["partner_country"].nunique())
    fit = sm.OLS(y_tilde, x_tilde, hasconst=False).fit(
        cov_type="cluster",
        cov_kwds={"groups": clusters, "use_correction": True},
        use_t=True,
    )
    beta = pd.Series(fit.params, index=terms)
    covariance = pd.DataFrame(fit.cov_params(), index=terms, columns=terms)

    pre_terms = [
        str(item["interaction_term"])
        for item in meta
        if int(item["event_time"]) <= -2
    ]
    pre_beta = beta.loc[pre_terms].to_numpy(dtype=float)
    pre_cov = covariance.loc[pre_terms, pre_terms].to_numpy(dtype=float)
    covariance_rank = int(np.linalg.matrix_rank(pre_cov))
    wald_stat = float(pre_beta @ np.linalg.pinv(pre_cov) @ pre_beta)
    pretrend_p = float(stats.chi2.sf(wald_stat, df=covariance_rank))

    critical = float(stats.t.ppf(0.975, df=cluster_count - 1))
    rows = []
    for event_time in sorted({int(item["event_time"]) for item in meta}):
        selected = [item for item in meta if int(item["event_time"]) == event_time]
        names = [str(item["interaction_term"]) for item in selected]
        raw_weights = np.array(
            [cohort_sizes[int(item["group"])] for item in selected], dtype=float
        )
        weights = raw_weights / raw_weights.sum()
        estimate = float(weights @ beta.loc[names].to_numpy(dtype=float))
        selected_cov = covariance.loc[names, names].to_numpy(dtype=float)
        se = float(np.sqrt(max(weights @ selected_cov @ weights, 0.0)))
        t_stat = estimate / se if se > 0 else np.nan
        p_value = (
            float(2 * stats.t.sf(abs(t_stat), df=cluster_count - 1))
            if se > 0
            else np.nan
        )
        rows.append(
            {
                "event_time": event_time,
                "revenue10_interaction": estimate,
                "std_error": se,
                "ci_low": estimate - critical * se,
                "ci_high": estimate + critical * se,
                "p_value_t18": p_value,
                "contributing_cohorts": ",".join(str(item["group"]) for item in selected),
            }
        )
    rows.append(
        {
            "event_time": -1,
            "revenue10_interaction": 0.0,
            "std_error": 0.0,
            "ci_low": 0.0,
            "ci_high": 0.0,
            "p_value_t18": np.nan,
            "contributing_cohorts": "reference",
        }
    )
    event = pd.DataFrame(rows).sort_values("event_time").reset_index(drop=True)
    maximum_absolute_pre = float(
        event.loc[event["event_time"].le(-2), "revenue10_interaction"].abs().max()
    )
    gate_pass = bool(pretrend_p >= 0.10 and maximum_absolute_pre <= 0.02)

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    event.to_csv(DATA / "event_study_gate.csv", index=False)
    term_table = pd.DataFrame(meta)
    term_table["main_estimate"] = term_table["main_term"].map(beta)
    term_table["interaction_estimate"] = term_table["interaction_term"].map(beta)
    term_table["interaction_std_error"] = term_table["interaction_term"].map(
        pd.Series(np.sqrt(np.diag(covariance)), index=terms)
    )
    term_table.to_csv(DATA / "cohort_event_terms.csv", index=False)
    export_table(
        event,
        "table2_event_study_gate",
        "Revenue-Dependence Moderation Event Study",
    )

    FIGURES.mkdir(parents=True, exist_ok=True)
    median = float(work["revenue_percent_pre"].median())
    work["high_revenue"] = work["revenue_percent_pre"].ge(median).astype("int8")
    raw = (
        work.groupby(["year", "rcep", "high_revenue"], as_index=False)["active"]
        .mean()
        .pivot(index=["year", "rcep"], columns="high_revenue", values="active")
        .reset_index()
    )
    raw["high_low_gap"] = raw[1] - raw[0]
    trends = raw.pivot(index="year", columns="rcep", values="high_low_gap")
    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.plot(trends.index, trends[0], marker="s", linestyle="--", color="#555555", label="Asian non-RCEP")
    ax.plot(trends.index, trends[1], marker="o", color="#0072B2", label="RCEP")
    ax.axvline(2021.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Year")
    ax.set_ylabel("High-minus-low quantified active rate")
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig1_raw_dependence_gap.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig1_raw_dependence_gap.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    plotted = event.loc[event["event_time"].ne(-1)]
    ax.errorbar(
        plotted["event_time"],
        plotted["revenue10_interaction"],
        yerr=[
            plotted["revenue10_interaction"] - plotted["ci_low"],
            plotted["ci_high"] - plotted["revenue10_interaction"],
        ],
        fmt="o-",
        color="#0072B2",
        ecolor="#555555",
        elinewidth=1,
        capsize=3,
        markersize=4,
        linewidth=1,
    )
    ax.scatter([-1], [0], facecolors="white", edgecolors="#0072B2", s=28, zorder=3)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.axvline(-0.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Years relative to RCEP entry")
    ax.set_ylabel("Effect difference per 10 revenue percentage points")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig2_event_study_gate.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig2_event_study_gate.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    diagnostics = {
        "experiment_id": "08_revenue_dependence",
        "gate": "REVENUE_MODERATION_PRETREND",
        "gate_pass": gate_pass,
        "joint_pretrend_wald_stat": wald_stat,
        "joint_pretrend_df": covariance_rank,
        "joint_pretrend_p_value": pretrend_p,
        "maximum_absolute_aggregated_pre_coefficient": maximum_absolute_pre,
        "economic_threshold": 0.02,
        "observations": int(len(work)),
        "relationships": int(work["relationship_id"].nunique()),
        "firm_year_fixed_effects": int(work["firm_year_id"].nunique()),
        "partner_country_clusters": cluster_count,
        "cohort_sizes": {str(key): int(value) for key, value in cohort_sizes.items()},
        "estimator": "OLS after pyhdfe residualization on pair-role and firm-year fixed effects; cohort-event main effects retained",
    }
    (LOGS / "pretrend_gate.json").write_text(json.dumps(diagnostics, indent=2), encoding="utf-8")
    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "PRETREND_GATE_RUN"
    status["results_inspected"] = True
    status["pretrend_gate_pass"] = gate_pass
    status["next_gate"] = "STATIC_MODERATION_AND_RANDOMIZATION" if gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(diagnostics, indent=2))


if __name__ == "__main__":
    main()
