#!/usr/bin/env python3
"""Build and gate the pre-policy quantified relationship panel."""

from __future__ import annotations

import hashlib
import json
import platform
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


EXPERIMENT = Path(__file__).resolve().parents[1]
REVISION = EXPERIMENT.parents[1]
WORKSPACE = REVISION.parent
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"
RCEP = {"JP", "KR", "AU", "NZ", "ID", "MY", "PH", "SG", "TH", "VN", "BN", "KH", "LA", "MM"}
LATE_2023 = {"ID", "PH"}
ASIA_CONTROLS = {"HK", "TW", "IN", "PK", "BD", "LK", "MN"}
COUNTRIES = RCEP | ASIA_CONTROLS
YEARS = list(range(2017, 2025))
CUTOFF = pd.Timestamp("2021-12-31")


def find_legacy_root() -> Path:
    for candidate in WORKSPACE.parent.iterdir():
        path = candidate / "data" / "cleaned" / "cn_global_links.parquet"
        if path.exists():
            return candidate.resolve()
    raise FileNotFoundError("Could not locate the verified legacy relationship source.")


def normalize_ticker(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    values = values.where(values.str.fullmatch(r"\d{1,6}"), pd.NA)
    return values.str.zfill(6)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = f"{value:.6f}" if isinstance(value, float) else str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def parse_end_dates(series: pd.Series) -> tuple[pd.Series, int]:
    text = series.astype("string").str.strip()
    open_marker = text.str.match(r"4000[-/]", na=False)
    parsed = pd.to_datetime(text.mask(open_marker), errors="coerce")
    parsed.loc[open_marker] = pd.Timestamp("2099-12-31")
    return parsed, int(open_marker.sum())


def main() -> None:
    legacy = find_legacy_root()
    source_path = legacy / "data" / "cleaned" / "cn_global_links.parquet"
    firm_path = legacy / "data" / "cleaned" / "ddd_panel.parquet"
    required = [
        "source_company_id", "target_company_id", "SOURCE_ticker", "TARGET_ticker",
        "cn_side", "rel_type", "partner_country", "start_dt", "end_dt",
        "revenue_percent",
    ]
    raw = pd.read_parquet(source_path, columns=required)
    firm = pd.read_parquet(firm_path, columns=["Stkcd", "year"])
    missing_columns = sorted(set(required) - set(raw.columns))
    if missing_columns:
        raise ValueError(f"Missing source columns: {missing_columns}")

    eligible_role = raw["rel_type"].isin(["CUSTOMER", "SUPPLIER"]) & raw["cn_side"].isin(["source", "target"])
    work = raw.loc[eligible_role].copy()
    work["partner_country"] = work["partner_country"].astype("string").str.strip().str.upper()
    work = work.loc[work["partner_country"].isin(COUNTRIES)].copy()
    source_is_cn = work["cn_side"].eq("source")
    work["listed_ticker"] = normalize_ticker(
        pd.Series(np.where(source_is_cn, work["SOURCE_ticker"], work["TARGET_ticker"]), index=work.index)
    )
    work["partner_id"] = np.where(
        source_is_cn, work["target_company_id"], work["source_company_id"]
    )
    work["supplier_role"] = (
        (source_is_cn & work["rel_type"].eq("SUPPLIER"))
        | (~source_is_cn & work["rel_type"].eq("CUSTOMER"))
    ).astype("int8")
    firm["listed_ticker"] = normalize_ticker(firm["Stkcd"])
    listed_tickers = set(firm["listed_ticker"].dropna())
    work = work.loc[
        work["listed_ticker"].isin(listed_tickers)
        & ~work["listed_ticker"].eq("000000")
        & work["partner_id"].notna()
    ].copy()
    work["start_date"] = pd.to_datetime(work["start_dt"], errors="coerce")
    work["end_date"], open_end_markers = parse_end_dates(work["end_dt"])
    invalid_dates = int((work["start_date"].isna() | work["end_date"].isna()).sum())
    work = work.dropna(subset=["start_date", "end_date"])
    work["revenue_percent_numeric"] = pd.to_numeric(work["revenue_percent"], errors="coerce")
    nonmissing_revenue = int(work["revenue_percent_numeric"].notna().sum())
    invalid_nonmissing_revenue = int(
        (
            work["revenue_percent_numeric"].notna()
            & (
                work["revenue_percent_numeric"].le(0)
                | work["revenue_percent_numeric"].gt(100)
            )
        ).sum()
    )
    relation_key = ["listed_ticker", "partner_id", "partner_country", "supplier_role"]
    pre_quantified = work.loc[
        work["start_date"].le(CUTOFF)
        & work["revenue_percent_numeric"].gt(0)
        & work["revenue_percent_numeric"].le(100)
    ].copy()
    exposure = (
        pre_quantified.groupby(relation_key, as_index=False)
        .agg(
            revenue_percent_pre=("revenue_percent_numeric", "max"),
            qualifying_pre_episodes=("start_date", "size"),
            first_pre_start=("start_date", "min"),
        )
    )
    episodes = work.merge(
        exposure[relation_key], on=relation_key, how="inner", validate="many_to_many"
    )
    episode_year = episodes.merge(pd.DataFrame({"year": YEARS}), how="cross")
    year_end = pd.to_datetime(episode_year["year"].astype(str) + "-12-31")
    episode_year["active_episode"] = (
        episode_year["start_date"].le(year_end) & episode_year["end_date"].ge(year_end)
    ).astype("int8")
    annual = (
        episode_year.groupby(relation_key + ["year"], as_index=False)
        .agg(active=("active_episode", "max"), episode_records=("active_episode", "size"))
    )
    panel = exposure.merge(annual, on=relation_key, how="left", validate="one_to_many")
    if panel["active"].isna().any():
        raise ValueError("Annual activity expansion produced missing values.")
    panel["rcep"] = panel["partner_country"].isin(RCEP).astype("int8")
    panel["entry_policy_year"] = np.where(
        panel["partner_country"].isin(LATE_2023),
        2023,
        np.where(panel["rcep"].eq(1), 2022, 9999),
    ).astype("int16")
    panel["treated_staged"] = (
        panel["rcep"].eq(1) & panel["year"].ge(panel["entry_policy_year"])
    ).astype("int8")
    panel["treated_uniform"] = (
        panel["rcep"].eq(1) & panel["year"].ge(2022)
    ).astype("int8")
    panel["event_time"] = np.where(
        panel["rcep"].eq(1), panel["year"] - panel["entry_policy_year"], np.nan
    )
    panel["revenue_percent10"] = panel["revenue_percent_pre"] / 10.0
    panel["treated_revenue_staged"] = panel["treated_staged"] * panel["revenue_percent10"]
    panel["treated_revenue_uniform"] = panel["treated_uniform"] * panel["revenue_percent10"]
    panel["relationship_id"] = pd.factorize(
        pd.MultiIndex.from_frame(panel[relation_key]), sort=True
    )[0].astype("int64")
    panel["firm_year_id"] = pd.factorize(
        pd.MultiIndex.from_frame(panel[["listed_ticker", "year"]]), sort=True
    )[0].astype("int64")

    if panel.duplicated(["relationship_id", "year"]).any():
        raise ValueError("Final relationship-year key is not unique.")
    balanced = panel.groupby("relationship_id")["year"].nunique().eq(len(YEARS)).all()
    relation = panel.loc[panel["year"].eq(2017)].copy()
    group_counts = relation.groupby("rcep")["relationship_id"].nunique()
    country_counts = relation.groupby("rcep")["partner_country"].nunique()
    firm_groups = relation.groupby("listed_ticker")["rcep"].agg(["min", "max"])
    overlap_firms = int((firm_groups["min"].eq(0) & firm_groups["max"].eq(1)).sum())
    within_variation = panel.groupby(["relationship_id", "rcep"])["active"].nunique().gt(1).groupby("rcep").sum()
    median_revenue = float(relation["revenue_percent_pre"].median())
    exposure_cells = (
        relation.assign(above_median=relation["revenue_percent_pre"].ge(median_revenue))
        .groupby(["rcep", "above_median"])["relationship_id"]
        .nunique()
    )
    total_relationships = int(relation["relationship_id"].nunique())
    total_countries = int(relation["partner_country"].nunique())
    pair_gate = (
        total_relationships >= 500
        and int(group_counts.get(0, 0)) >= 150
        and int(group_counts.get(1, 0)) >= 150
    )
    country_gate = (
        total_countries >= 10
        and int(country_counts.get(0, 0)) >= 4
        and int(country_counts.get(1, 0)) >= 4
    )
    firm_gate = overlap_firms >= 30
    variation_gate = all(int(within_variation.get(group, 0)) >= 100 for group in [0, 1])
    exposure_gate = all(
        int(exposure_cells.get((group, above), 0)) >= 50
        for group in [0, 1]
        for above in [False, True]
    )
    gate_pass = bool(
        pair_gate and country_gate and firm_gate and variation_gate and exposure_gate and balanced
    )

    support_table = pd.DataFrame(
        [
            {"metric": "All quantified pair-roles", "value": total_relationships, "threshold": 500, "pass": total_relationships >= 500},
            {"metric": "Control quantified pair-roles", "value": int(group_counts.get(0, 0)), "threshold": 150, "pass": int(group_counts.get(0, 0)) >= 150},
            {"metric": "RCEP quantified pair-roles", "value": int(group_counts.get(1, 0)), "threshold": 150, "pass": int(group_counts.get(1, 0)) >= 150},
            {"metric": "All countries", "value": total_countries, "threshold": 10, "pass": total_countries >= 10},
            {"metric": "Control countries", "value": int(country_counts.get(0, 0)), "threshold": 4, "pass": int(country_counts.get(0, 0)) >= 4},
            {"metric": "RCEP countries", "value": int(country_counts.get(1, 0)), "threshold": 4, "pass": int(country_counts.get(1, 0)) >= 4},
            {"metric": "Firms with both groups", "value": overlap_firms, "threshold": 30, "pass": overlap_firms >= 30},
            {"metric": "Control pairs with active variation", "value": int(within_variation.get(0, 0)), "threshold": 100, "pass": int(within_variation.get(0, 0)) >= 100},
            {"metric": "RCEP pairs with active variation", "value": int(within_variation.get(1, 0)), "threshold": 100, "pass": int(within_variation.get(1, 0)) >= 100},
            {"metric": "Minimum group-by-median exposure cell", "value": min(int(exposure_cells.get((group, above), 0)) for group in [0, 1] for above in [False, True]), "threshold": 50, "pass": exposure_gate},
        ]
    )
    revenue_summary = relation.groupby("rcep", as_index=False)["revenue_percent_pre"].agg(
        count="count", mean="mean", std="std", minimum="min", median="median", maximum="max"
    )

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    panel = panel.sort_values(["relationship_id", "year"]).reset_index(drop=True)
    panel.to_parquet(DATA / "quantified_relationship_year.parquet", index=False)
    revenue_summary.to_csv(DATA / "revenue_percent_summary.csv", index=False)
    export_table(support_table, "table1_support_gate", "Quantified Relationship Support Gate")
    gate = {
        "experiment_id": "08_revenue_dependence",
        "gate": "DATA_SUPPORT",
        "gate_pass": gate_pass,
        "pair_support_gate": bool(pair_gate),
        "country_support_gate": bool(country_gate),
        "within_firm_support_gate": bool(firm_gate),
        "outcome_variation_gate": bool(variation_gate),
        "exposure_support_gate": bool(exposure_gate),
        "balanced_panel_gate": bool(balanced),
        "quantified_pair_roles": total_relationships,
        "control_pair_roles": int(group_counts.get(0, 0)),
        "rcep_pair_roles": int(group_counts.get(1, 0)),
        "countries": total_countries,
        "control_countries": int(country_counts.get(0, 0)),
        "rcep_countries": int(country_counts.get(1, 0)),
        "firms_with_both_groups": overlap_firms,
        "control_pairs_with_active_variation": int(within_variation.get(0, 0)),
        "rcep_pairs_with_active_variation": int(within_variation.get(1, 0)),
        "median_revenue_percent": median_revenue,
        "minimum_group_by_median_cell": min(int(exposure_cells.get((group, above), 0)) for group in [0, 1] for above in [False, True]),
    }
    (DATA / "support_gate.json").write_text(json.dumps(gate, indent=2), encoding="utf-8")
    summary = {
        "experiment_id": "08_revenue_dependence",
        "build_status": "PASSED",
        "python": platform.python_version(),
        "source_hashes": {
            "episode_source_sha256": sha256(source_path),
            "firm_source_sha256": sha256(firm_path),
        },
        "source_rows": int(len(raw)),
        "eligible_role_country_listed_rows": int(len(work)),
        "open_end_markers_recoded": open_end_markers,
        "invalid_date_rows_excluded": invalid_dates,
        "nonmissing_revenue_rows": nonmissing_revenue,
        "invalid_nonmissing_revenue_rows": invalid_nonmissing_revenue,
        "qualifying_prepolicy_quantified_rows": int(len(pre_quantified)),
        "support_gate": gate,
    }
    (DATA / "panel_build_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (LOGS / "panel_build.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    sample_path = DATA / "sample_construction.json"
    sample = json.loads(sample_path.read_text(encoding="utf-8"))
    sample["status"] = "COMPLETED"
    sample["counts"] = summary
    sample["execution_log"] = "logs/panel_build.json"
    sample_path.write_text(json.dumps(sample, indent=2), encoding="utf-8")
    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "DATA_SUPPORT_GATE_RUN"
    status["results_inspected"] = True
    status["support_gate_pass"] = gate_pass
    status["next_gate"] = "PRETREND" if gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
