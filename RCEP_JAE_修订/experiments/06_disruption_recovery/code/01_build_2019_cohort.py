#!/usr/bin/env python3
"""Build and gate the fixed-2019 pre-pandemic supplier-link cohort."""

from __future__ import annotations

import hashlib
import json
import platform
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


EXPERIMENT = Path(__file__).resolve().parents[1]
REVISION = EXPERIMENT.parents[1]
WORKSPACE = REVISION.parent
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"
RCEP = {"JP", "KR", "AU", "NZ", "ID", "MY", "PH", "SG", "TH", "VN", "BN", "KH", "LA", "MM"}
LATE_2023 = {"ID", "PH"}
ASIA_CONTROLS = {"HK", "TW", "IN", "PK", "BD", "LK", "MN"}
COUNTRIES = RCEP | ASIA_CONTROLS
YEARS = list(range(2017, 2025))
FOLLOW_YEARS = list(range(2020, 2025))


def find_legacy_root() -> Path:
    for candidate in WORKSPACE.parent.iterdir():
        path = candidate / "data" / "cleaned" / "ddd_panel.parquet"
        if path.exists():
            return candidate.resolve()
    raise FileNotFoundError("Could not locate the verified legacy data root.")


def normalize_ticker(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    values = values.where(values.str.fullmatch(r"\d{1,6}"), pd.NA)
    return values.str.zfill(6)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = f"{value:.6f}" if isinstance(value, float) else str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    pair_path = REVISION / "data" / "derived" / "pair_year.parquet"
    firm_path = find_legacy_root() / "data" / "cleaned" / "ddd_panel.parquet"
    pair = pd.read_parquet(
        pair_path,
        columns=[
            "cn_id", "partner_id", "partner_country", "supplier_link", "pair_id",
            "year", "active", "ticker_key",
        ],
    )
    firm = pd.read_parquet(firm_path, columns=["Stkcd", "year"])
    duplicate_source = int(pair.duplicated(["pair_id", "year"]).sum())
    invalid_active = int((~pair["active"].isin([0, 1])).sum())
    invalid_year = int((~pair["year"].isin(YEARS)).sum())
    if duplicate_source or invalid_active or invalid_year:
        raise ValueError(
            f"Source validation failed: duplicate={duplicate_source}, "
            f"invalid_active={invalid_active}, invalid_year={invalid_year}"
        )

    pair["listed_ticker"] = normalize_ticker(pair["ticker_key"])
    firm["listed_ticker"] = normalize_ticker(firm["Stkcd"])
    listed_tickers = set(firm["listed_ticker"].dropna())
    invalid_or_zero_rows = int(
        (pair["listed_ticker"].isna() | pair["listed_ticker"].eq("000000")).sum()
    )
    nonlisted_rows = int(
        (
            pair["listed_ticker"].notna()
            & ~pair["listed_ticker"].eq("000000")
            & ~pair["listed_ticker"].isin(listed_tickers)
        ).sum()
    )
    eligible = pair.loc[
        pair["supplier_link"].eq(1)
        & pair["partner_country"].isin(COUNTRIES)
        & pair["listed_ticker"].isin(listed_tickers)
        & ~pair["listed_ticker"].eq("000000")
    ].copy()
    relation_key = ["listed_ticker", "partner_id", "partner_country"]
    collapsed = (
        eligible.groupby(relation_key + ["year"], as_index=False)
        .agg(current_active=("active", "max"), factset_cn_ids=("cn_id", "nunique"))
    )
    cohort_keys = collapsed.loc[
        collapsed["year"].eq(2019) & collapsed["current_active"].eq(1), relation_key
    ].drop_duplicates()
    cohort = collapsed.merge(cohort_keys, on=relation_key, how="inner", validate="many_to_one")
    histories = cohort.groupby(relation_key)["year"].nunique()
    incomplete_histories = int(histories.ne(len(YEARS)).sum())
    if incomplete_histories:
        raise ValueError(f"Cohort has {incomplete_histories} incomplete histories.")
    cohort = cohort.sort_values(relation_key + ["year"]).reset_index(drop=True)
    cohort["survival"] = (
        cohort.loc[cohort["year"].ge(2019)]
        .groupby(relation_key)["current_active"]
        .cummin()
        .reindex(cohort.index)
    )
    cohort["lag_active"] = cohort.groupby(relation_key)["current_active"].shift(1)
    cohort["reactivation"] = (
        cohort["current_active"].eq(1) & cohort["lag_active"].eq(0)
    ).astype("int8")
    cohort = cohort.loc[cohort["year"].isin(FOLLOW_YEARS)].copy()
    cohort["survival"] = cohort["survival"].astype("int8")
    cohort["lag_active"] = cohort["lag_active"].astype("int8")
    cohort["rcep"] = cohort["partner_country"].isin(RCEP).astype("int8")
    cohort["entry_policy_year"] = np.where(
        cohort["partner_country"].isin(LATE_2023),
        2023,
        np.where(cohort["rcep"].eq(1), 2022, 9999),
    ).astype("int16")
    cohort["treated_staged"] = (
        cohort["rcep"].eq(1) & cohort["year"].ge(cohort["entry_policy_year"])
    ).astype("int8")
    cohort["treated_uniform"] = (
        cohort["rcep"].eq(1) & cohort["year"].ge(2022)
    ).astype("int8")
    cohort["event_time"] = np.where(
        cohort["rcep"].eq(1), cohort["year"] - cohort["entry_policy_year"], np.nan
    )
    cohort["relationship_id"] = pd.factorize(
        pd.MultiIndex.from_frame(cohort[relation_key]), sort=True
    )[0].astype("int64")
    cohort["firm_year_id"] = pd.factorize(
        pd.MultiIndex.from_frame(cohort[["listed_ticker", "year"]]), sort=True
    )[0].astype("int64")

    if cohort.duplicated(relation_key + ["year"]).any():
        raise ValueError("Final cohort key is not unique.")
    balanced = cohort.groupby("relationship_id")["year"].nunique().eq(len(FOLLOW_YEARS)).all()
    baseline = cohort.loc[cohort["year"].eq(2020)].copy()
    group_counts = baseline.groupby("rcep")["relationship_id"].nunique()
    positive_countries = baseline.groupby("rcep")["partner_country"].nunique()
    country_counts = baseline.groupby(["partner_country", "rcep"], as_index=False).agg(
        relationships=("relationship_id", "nunique"),
        listed_firms=("listed_ticker", "nunique"),
    )
    firm_groups = baseline.groupby("listed_ticker")["rcep"].agg(["min", "max"])
    overlap_firms = int((firm_groups["min"].eq(0) & firm_groups["max"].eq(1)).sum())
    total_relationships = int(baseline["relationship_id"].nunique())
    cluster_count = int(baseline["partner_country"].nunique())
    cohort_gate = bool(
        total_relationships >= 500
        and int(group_counts.get(0, 0)) >= 150
        and int(group_counts.get(1, 0)) >= 150
    )
    country_gate = bool(
        int(positive_countries.get(0, 0)) >= 5
        and int(positive_countries.get(1, 0)) >= 5
        and cluster_count >= 12
    )
    firm_gate = overlap_firms >= 50
    gate_pass = bool(cohort_gate and country_gate and firm_gate and balanced)

    year_support = cohort.groupby(["year", "rcep"], as_index=False).agg(
        relationships=("relationship_id", "nunique"),
        listed_firms=("listed_ticker", "nunique"),
        active_links=("current_active", "sum"),
        surviving_links=("survival", "sum"),
        reactivations=("reactivation", "sum"),
    )
    support_table = pd.DataFrame(
        [
            {"metric": "Total cohort relationships", "value": total_relationships, "threshold": 500, "pass": total_relationships >= 500},
            {"metric": "Control cohort relationships", "value": int(group_counts.get(0, 0)), "threshold": 150, "pass": int(group_counts.get(0, 0)) >= 150},
            {"metric": "RCEP cohort relationships", "value": int(group_counts.get(1, 0)), "threshold": 150, "pass": int(group_counts.get(1, 0)) >= 150},
            {"metric": "Control countries", "value": int(positive_countries.get(0, 0)), "threshold": 5, "pass": int(positive_countries.get(0, 0)) >= 5},
            {"metric": "RCEP countries", "value": int(positive_countries.get(1, 0)), "threshold": 5, "pass": int(positive_countries.get(1, 0)) >= 5},
            {"metric": "All country clusters", "value": cluster_count, "threshold": 12, "pass": cluster_count >= 12},
            {"metric": "Firms with both groups", "value": overlap_firms, "threshold": 50, "pass": overlap_firms >= 50},
            {"metric": "Balanced five-year panel", "value": int(balanced), "threshold": 1, "pass": bool(balanced)},
        ]
    )

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    cohort = cohort.sort_values(["relationship_id", "year"]).reset_index(drop=True)
    cohort.to_parquet(DATA / "prepandemic_relationship_cohort.parquet", index=False)
    country_counts.to_csv(DATA / "country_cohort_support.csv", index=False)
    year_support.to_csv(DATA / "year_group_support.csv", index=False)
    export_table(support_table, "table1_support_gate", "Pre-pandemic Cohort Support Gate")
    gate = {
        "experiment_id": "06_disruption_recovery",
        "gate": "DATA_SUPPORT",
        "gate_pass": gate_pass,
        "cohort_size_gate": cohort_gate,
        "country_support_gate": country_gate,
        "within_firm_support_gate": firm_gate,
        "balanced_panel_gate": bool(balanced),
        "total_relationships": total_relationships,
        "control_relationships": int(group_counts.get(0, 0)),
        "rcep_relationships": int(group_counts.get(1, 0)),
        "control_countries": int(positive_countries.get(0, 0)),
        "rcep_countries": int(positive_countries.get(1, 0)),
        "country_clusters": cluster_count,
        "firms_with_both_groups": overlap_firms,
    }
    (DATA / "support_gate.json").write_text(json.dumps(gate, indent=2), encoding="utf-8")
    summary = {
        "experiment_id": "06_disruption_recovery",
        "build_status": "PASSED",
        "python": platform.python_version(),
        "source_hashes": {
            "pair_year_sha256": sha256(pair_path),
            "firm_year_sha256": sha256(firm_path),
        },
        "source_rows": int(len(pair)),
        "invalid_or_zero_ticker_rows": invalid_or_zero_rows,
        "nonlisted_ticker_rows": nonlisted_rows,
        "eligible_rows": int(len(eligible)),
        "collapsed_rows": int(len(collapsed)),
        "support_gate": gate,
    }
    (DATA / "cohort_build_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (LOGS / "cohort_build.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    sample_path = DATA / "sample_construction.json"
    sample = json.loads(sample_path.read_text(encoding="utf-8"))
    sample["status"] = "COMPLETED"
    sample["counts"] = summary
    sample["execution_log"] = "logs/cohort_build.json"
    sample_path.write_text(json.dumps(sample, indent=2), encoding="utf-8")
    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "DATA_SUPPORT_GATE_RUN"
    status["results_inspected"] = True
    status["support_gate_pass"] = gate_pass
    status["next_gate"] = "PRETREND" if gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
