#!/usr/bin/env python3
"""Build and gate the frozen firm-country first-entry risk set."""

from __future__ import annotations

import hashlib
import json
import platform
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


EXPERIMENT = Path(__file__).resolve().parents[1]
REVISION = EXPERIMENT.parents[1]
WORKSPACE = REVISION.parent
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"
RCEP = {"JP", "KR", "AU", "NZ", "ID", "MY", "PH", "SG", "TH", "VN", "BN", "KH", "LA", "MM"}
LATE_2023 = {"ID", "PH"}
ASIA_CONTROLS = {"HK", "TW", "IN", "PK", "BD", "LK", "MN"}
COUNTRIES = sorted(RCEP | ASIA_CONTROLS)
YEARS = list(range(2017, 2025))
ANALYSIS_YEARS = list(range(2018, 2025))


def find_legacy_root() -> Path:
    for candidate in WORKSPACE.parent.iterdir():
        ddd = candidate / "data" / "cleaned" / "ddd_panel.parquet"
        if ddd.exists():
            return candidate.resolve()
    raise FileNotFoundError("Could not locate the verified legacy data root.")


def normalize_ticker(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    values = values.where(values.str.fullmatch(r"\d{1,6}"), pd.NA)
    return values.str.zfill(6)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = f"{value:.6f}" if isinstance(value, float) else str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    pair_path = REVISION / "data" / "derived" / "pair_year.parquet"
    firm_path = find_legacy_root() / "data" / "cleaned" / "ddd_panel.parquet"
    pair = pd.read_parquet(
        pair_path,
        columns=[
            "cn_id", "partner_id", "partner_country", "supplier_link", "pair_id",
            "year", "active", "first_start", "ticker_key",
        ],
    )
    firm = pd.read_parquet(firm_path, columns=["Stkcd", "year"])

    duplicate_pair_year = int(pair.duplicated(["pair_id", "year"]).sum())
    invalid_active = int((~pair["active"].isin([0, 1])).sum())
    invalid_year = int((~pair["year"].isin(YEARS)).sum())
    if duplicate_pair_year or invalid_active or invalid_year:
        raise ValueError(
            f"Source validation failed: duplicate={duplicate_pair_year}, "
            f"invalid_active={invalid_active}, invalid_year={invalid_year}"
        )

    pair["listed_ticker"] = normalize_ticker(pair["ticker_key"])
    firm["listed_ticker"] = normalize_ticker(firm["Stkcd"])
    listed_tickers = set(firm["listed_ticker"].dropna())
    mapped = pair.loc[
        pair["listed_ticker"].isin(listed_tickers)
        & ~pair["listed_ticker"].eq("000000")
    ].copy()
    supplier = mapped.loc[mapped["supplier_link"].eq(1)].copy()

    covered_firms = sorted(
        supplier.loc[
            supplier["year"].eq(2017)
            & supplier["active"].eq(1)
            & supplier["partner_country"].astype("string").str.fullmatch(r"[A-Z]{2}", na=False),
            "listed_ticker",
        ].unique()
    )
    if not covered_firms:
        raise ValueError("No listed firms meet the frozen 2017 supplier-coverage rule.")

    observed_country = (
        supplier.loc[supplier["partner_country"].isin(COUNTRIES)]
        .groupby(["listed_ticker", "partner_country", "year"], as_index=False)
        .agg(current_active=("active", "max"))
    )
    first_start = (
        supplier.loc[supplier["partner_country"].isin(COUNTRIES)]
        .groupby(["listed_ticker", "partner_country"], as_index=False)
        .agg(first_supplier_start=("first_start", "min"))
    )

    dyads = pd.MultiIndex.from_product(
        [covered_firms, COUNTRIES], names=["listed_ticker", "partner_country"]
    ).to_frame(index=False)
    dyads = dyads.merge(
        first_start,
        on=["listed_ticker", "partner_country"],
        how="left",
        validate="one_to_one",
    )
    dyads["preexisting_by_2017"] = dyads["first_supplier_start"].le(
        pd.Timestamp("2017-12-31")
    )
    preexisting_dyads = int(dyads["preexisting_by_2017"].sum())
    eligible_dyads = dyads.loc[~dyads["preexisting_by_2017"]].copy()

    risk = eligible_dyads.merge(pd.DataFrame({"year": ANALYSIS_YEARS}), how="cross")
    risk = risk.merge(
        observed_country,
        on=["listed_ticker", "partner_country", "year"],
        how="left",
        validate="one_to_one",
    )
    risk["current_active"] = risk["current_active"].fillna(0).astype("int8")
    first_entry = (
        risk.loc[risk["current_active"].eq(1)]
        .groupby(["listed_ticker", "partner_country"])["year"]
        .min()
        .rename("first_entry_year")
    )
    risk = risk.merge(
        first_entry,
        on=["listed_ticker", "partner_country"],
        how="left",
        validate="many_to_one",
    )
    risk = risk.loc[
        risk["first_entry_year"].isna() | risk["year"].le(risk["first_entry_year"])
    ].copy()
    risk["first_country_entry"] = risk["year"].eq(risk["first_entry_year"]).astype("int8")

    partner = (
        supplier.loc[supplier["partner_country"].isin(COUNTRIES)]
        .groupby(
            ["listed_ticker", "partner_country", "partner_id", "year"], as_index=False
        )["active"]
        .max()
    )
    partner_first_active = (
        partner.loc[partner["active"].eq(1)]
        .groupby(["listed_ticker", "partner_country", "partner_id"])["year"]
        .min()
        .rename("partner_first_active_year")
        .reset_index()
    )
    new_partner_counts = (
        partner_first_active.loc[partner_first_active["partner_first_active_year"].ge(2018)]
        .groupby(
            ["listed_ticker", "partner_country", "partner_first_active_year"], as_index=False
        )
        .size()
        .rename(
            columns={
                "partner_first_active_year": "year",
                "size": "first_partner_link_entries",
            }
        )
    )
    risk = risk.merge(
        new_partner_counts,
        on=["listed_ticker", "partner_country", "year"],
        how="left",
        validate="one_to_one",
    )
    risk["first_partner_link_entries"] = (
        risk["first_partner_link_entries"].fillna(0).astype("int16")
    )

    risk["rcep"] = risk["partner_country"].isin(RCEP).astype("int8")
    risk["entry_policy_year"] = np.where(
        risk["partner_country"].isin(LATE_2023),
        2023,
        np.where(risk["rcep"].eq(1), 2022, 9999),
    ).astype("int16")
    risk["treated_staged"] = (
        risk["rcep"].eq(1) & risk["year"].ge(risk["entry_policy_year"])
    ).astype("int8")
    risk["treated_uniform"] = (
        risk["rcep"].eq(1) & risk["year"].ge(2022)
    ).astype("int8")
    risk["event_time"] = np.where(
        risk["rcep"].eq(1), risk["year"] - risk["entry_policy_year"], np.nan
    )
    risk["period"] = np.where(risk["year"].le(2021), "pre", "post")
    risk["firm_country_id"] = pd.factorize(
        pd.MultiIndex.from_frame(risk[["listed_ticker", "partner_country"]]), sort=True
    )[0].astype("int64")
    risk["firm_year_id"] = pd.factorize(
        pd.MultiIndex.from_frame(risk[["listed_ticker", "year"]]), sort=True
    )[0].astype("int64")

    if risk.duplicated(["listed_ticker", "partner_country", "year"]).any():
        raise ValueError("Risk-set key is not unique.")
    entry_per_dyad = risk.groupby(["listed_ticker", "partner_country"])[
        "first_country_entry"
    ].sum()
    if entry_per_dyad.gt(1).any():
        raise ValueError("At least one firm-country dyad enters more than once.")
    entered = risk.loc[risk["first_country_entry"].eq(1), ["listed_ticker", "partner_country", "year"]]
    if not entered.empty:
        check = risk.merge(
            entered.rename(columns={"year": "entered_year"}),
            on=["listed_ticker", "partner_country"],
            how="inner",
        )
        if check["year"].gt(check["entered_year"]).any():
            raise ValueError("Rows remain after first country entry.")

    support = (
        risk.groupby(["rcep", "period"], as_index=False)
        .agg(
            risk_rows=("first_country_entry", "size"),
            firms=("listed_ticker", "nunique"),
            countries=("partner_country", "nunique"),
            first_country_entries=("first_country_entry", "sum"),
            first_partner_link_entries=("first_partner_link_entries", "sum"),
        )
    )
    support["entry_rate"] = support["first_country_entries"] / support["risk_rows"]
    country_period = (
        risk.groupby(["partner_country", "rcep", "period"], as_index=False)
        .agg(
            risk_rows=("first_country_entry", "size"),
            firms=("listed_ticker", "nunique"),
            first_country_entries=("first_country_entry", "sum"),
        )
    )
    year_support = (
        risk.groupby(["year", "rcep"], as_index=False)
        .agg(
            risk_rows=("first_country_entry", "size"),
            firms=("listed_ticker", "nunique"),
            countries=("partner_country", "nunique"),
            first_country_entries=("first_country_entry", "sum"),
        )
    )
    firm_year_group = (
        risk.groupby(["listed_ticker", "year", "rcep"], as_index=False)
        .size()
        .pivot_table(
            index=["listed_ticker", "year"], columns="rcep", values="size", fill_value=0
        )
        .reset_index()
    )
    for column in [0, 1]:
        if column not in firm_year_group:
            firm_year_group[column] = 0
    overlap_by_year = (
        firm_year_group.assign(has_both=firm_year_group[0].gt(0) & firm_year_group[1].gt(0))
        .groupby("year", as_index=False)
        .agg(firms_with_both_groups=("has_both", "sum"))
    )

    support_key = support.set_index(["rcep", "period"])["first_country_entries"]
    required_keys = [(0, "pre"), (1, "pre"), (0, "post"), (1, "post")]
    event_gate = all(int(support_key.get(key, 0)) >= 30 for key in required_keys)
    event_countries = country_period.loc[country_period["first_country_entries"].gt(0)].groupby(
        ["rcep", "period"]
    )["partner_country"].nunique()
    country_gate = all(int(event_countries.get(key, 0)) >= 5 for key in required_keys)
    overlap_gate = bool(overlap_by_year["firms_with_both_groups"].ge(100).all())
    support_gate_pass = bool(event_gate and country_gate and overlap_gate)

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    risk = risk.sort_values(["firm_country_id", "year"]).reset_index(drop=True)
    risk.to_parquet(DATA / "country_entry_risk_set.parquet", index=False)
    support.to_csv(DATA / "support_by_group_period.csv", index=False)
    country_period.to_csv(DATA / "country_period_support.csv", index=False)
    year_support.to_csv(DATA / "year_group_support.csv", index=False)
    overlap_by_year.to_csv(DATA / "within_firm_year_overlap.csv", index=False)
    export_table(support, "table1_support_gate", "Country Entry Risk-Set Support")

    gate = {
        "experiment_id": "03_new_link_formation",
        "gate": "DATA_SUPPORT",
        "gate_pass": support_gate_pass,
        "event_count_gate": event_gate,
        "event_country_gate": country_gate,
        "within_firm_year_overlap_gate": overlap_gate,
        "minimum_events_per_group_period": 30,
        "minimum_event_countries_per_group_period": 5,
        "minimum_overlapping_firms_per_year": 100,
        "event_counts": {
            f"rcep_{key[0]}_{key[1]}": int(support_key.get(key, 0)) for key in required_keys
        },
        "event_country_counts": {
            f"rcep_{key[0]}_{key[1]}": int(event_countries.get(key, 0)) for key in required_keys
        },
        "minimum_observed_overlap_firms": int(
            overlap_by_year["firms_with_both_groups"].min()
        ),
    }
    (DATA / "support_gate.json").write_text(json.dumps(gate, indent=2), encoding="utf-8")

    summary = {
        "experiment_id": "03_new_link_formation",
        "build_status": "PASSED",
        "python": platform.python_version(),
        "source_hashes": {
            "pair_year_sha256": sha256(pair_path),
            "firm_year_sha256": sha256(firm_path),
        },
        "source_pair_year_rows": int(len(pair)),
        "mapped_pair_year_rows": int(len(mapped)),
        "frozen_2017_covered_firms": int(len(covered_firms)),
        "candidate_country_dyads_before_history_exclusion": int(len(dyads)),
        "preexisting_country_dyads_excluded": preexisting_dyads,
        "eligible_country_dyads": int(len(eligible_dyads)),
        "risk_set_rows": int(len(risk)),
        "first_country_entries": int(risk["first_country_entry"].sum()),
        "support_gate": gate,
        "outputs": [
            "data/country_entry_risk_set.parquet",
            "data/support_by_group_period.csv",
            "data/country_period_support.csv",
            "data/year_group_support.csv",
            "data/within_firm_year_overlap.csv",
            "tables/table1_support_gate.xlsx",
            "tables/table1_support_gate.docx",
            "tables/table1_support_gate.tex"
        ]
    }
    (DATA / "risk_set_build_summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    (LOGS / "risk_set_build.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    sample_path = DATA / "sample_construction.json"
    sample = json.loads(sample_path.read_text(encoding="utf-8"))
    sample["status"] = "COMPLETED"
    sample["counts"] = summary
    sample["execution_log"] = "logs/risk_set_build.json"
    sample_path.write_text(json.dumps(sample, indent=2), encoding="utf-8")

    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "DATA_SUPPORT_GATE_RUN"
    status["results_inspected"] = True
    status["support_gate_pass"] = support_gate_pass
    status["next_gate"] = "PRETREND_EVENT_STUDY" if support_gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
