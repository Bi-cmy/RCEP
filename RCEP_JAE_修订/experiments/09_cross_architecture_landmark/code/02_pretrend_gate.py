#!/usr/bin/env python3
"""Run the frozen one-lead landmark-design pretrend gate."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from pyhdfe import create
from scipy import stats


EXPERIMENT = Path(__file__).resolve().parents[1]
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
FIGURES = EXPERIMENT / "figures"
LOGS = EXPERIMENT / "logs"
STATUS = EXPERIMENT / "STATUS.json"


plt.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": ["Times New Roman", "DejaVu Serif"],
        "font.size": 9,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 300,
    }
)


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, float):
                cell.text = "" if np.isnan(value) else f"{value:.6f}"
            else:
                cell.text = str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    panel = pd.read_parquet(DATA / "architecture_landmark_panel.parquet")
    support = json.loads((DATA / "support_replication.json").read_text(encoding="utf-8"))
    if not support["gate_pass"]:
        raise RuntimeError("Support replication did not pass.")
    if panel.duplicated(["listed_ticker", "year"]).any():
        raise ValueError("Firm-year key is not unique.")

    work = panel.copy()
    event_years = [2020, 2022, 2023, 2024]
    terms = []
    for year in event_years:
        term = f"architecture_x_{year}"
        work[term] = work["architecture_count"] * work["year"].eq(year)
        terms.append(term)
    y = work["asinh_active_rcep_supplier_links"].to_numpy(dtype=float)
    x = work[terms].to_numpy(dtype=float)
    fixed_effects = work[["firm_id", "industry_year_id"]].to_numpy(dtype=np.int64)
    algorithm = create(fixed_effects, drop_singletons=False)
    transformed = algorithm.residualize(np.column_stack([y, x]))
    y_tilde, x_tilde = transformed[:, 0], transformed[:, 1:]
    clusters = work["firm_id"].to_numpy(dtype=np.int64)
    cluster_count = int(work["firm_id"].nunique())
    fit = sm.OLS(y_tilde, x_tilde, hasconst=False).fit(
        cov_type="cluster",
        cov_kwds={"groups": clusters, "use_correction": True},
        use_t=True,
    )
    critical = float(stats.t.ppf(0.975, df=cluster_count - 1))
    event_rows = [
        {
            "year": year,
            "relative_to_2021": year - 2021,
            "estimate_per_architecture": float(fit.params[index]),
            "std_error": float(fit.bse[index]),
            "ci_low": float(fit.params[index] - critical * fit.bse[index]),
            "ci_high": float(fit.params[index] + critical * fit.bse[index]),
            "p_value": float(fit.pvalues[index]),
        }
        for index, year in enumerate(event_years)
    ]
    event_rows.append(
        {
            "year": 2021,
            "relative_to_2021": 0,
            "estimate_per_architecture": 0.0,
            "std_error": 0.0,
            "ci_low": 0.0,
            "ci_high": 0.0,
            "p_value": np.nan,
        }
    )
    event = pd.DataFrame(event_rows).sort_values("year").reset_index(drop=True)
    lead = event.loc[event["year"].eq(2020)].iloc[0]
    lead_estimate = float(lead["estimate_per_architecture"])
    lead_p = float(lead["p_value"])
    gate_pass = bool(lead_p >= 0.10 and abs(lead_estimate) <= 0.05)

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    event.to_csv(DATA / "event_study_gate.csv", index=False)
    export_table(event, "table1_event_study_gate", "Cross-Architecture Landmark Event Study")

    FIGURES.mkdir(parents=True, exist_ok=True)
    work["any_architecture"] = work["architecture_count"].ge(1).astype("int8")
    trends = work.groupby(["year", "any_architecture"])["asinh_active_rcep_supplier_links"].mean().unstack()
    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.plot(trends.index, trends[0], marker="s", linestyle="--", color="#555555", label="No frozen architecture")
    ax.plot(trends.index, trends[1], marker="o", color="#0072B2", label="One or more architectures")
    ax.axvline(2021.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Year")
    ax.set_ylabel("Mean asinh active RCEP supplier links")
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig1_raw_trends.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig1_raw_trends.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    plotted = event.loc[event["year"].ne(2021)]
    ax.errorbar(
        plotted["relative_to_2021"],
        plotted["estimate_per_architecture"],
        yerr=[
            plotted["estimate_per_architecture"] - plotted["ci_low"],
            plotted["ci_high"] - plotted["estimate_per_architecture"],
        ],
        fmt="o-",
        color="#0072B2",
        ecolor="#555555",
        elinewidth=1,
        capsize=3,
        markersize=4,
        linewidth=1,
    )
    ax.scatter([0], [0], facecolors="white", edgecolors="#0072B2", s=28, zorder=3)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.axvline(0.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Year relative to 2021")
    ax.set_ylabel("Difference per frozen architecture")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig2_event_study_gate.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig2_event_study_gate.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    diagnostics = {
        "experiment_id": "09_cross_architecture_landmark",
        "gate": "ONE_LEAD_PRETREND",
        "gate_pass": gate_pass,
        "lead_year": 2020,
        "reference_year": 2021,
        "lead_estimate_per_architecture": lead_estimate,
        "lead_std_error": float(lead["std_error"]),
        "lead_p_value": lead_p,
        "economic_threshold": 0.05,
        "observations": int(len(work)),
        "firms": cluster_count,
        "industry_year_fixed_effects": int(work["industry_year_id"].nunique()),
        "estimator": "OLS after pyhdfe residualization on firm and frozen-industry-year fixed effects",
        "diagnostic_limit": "Only one clean non-overlapping pre-policy interaction is available.",
    }
    (LOGS / "pretrend_gate.json").write_text(json.dumps(diagnostics, indent=2), encoding="utf-8")
    status = json.loads(STATUS.read_text(encoding="utf-8"))
    status["phase"] = "PRETREND_GATE_RUN"
    status["pretrend_gate_pass"] = gate_pass
    status["next_gate"] = "STATIC_AND_STRATIFIED_RANDOMIZATION" if gate_pass else "ARCHIVE_FAILURE"
    STATUS.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(diagnostics, indent=2))


if __name__ == "__main__":
    main()
