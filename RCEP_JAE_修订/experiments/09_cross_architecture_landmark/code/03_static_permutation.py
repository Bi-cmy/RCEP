#!/usr/bin/env python3
"""Estimate the landmark interaction and its frozen stratified inference checks."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from pyhdfe import create


EXPERIMENT = Path(__file__).resolve().parents[1]
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
FIGURES = EXPERIMENT / "figures"
LOGS = EXPERIMENT / "logs"
STATUS = EXPERIMENT / "STATUS.json"
SEED = 20260723
B = 999


plt.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": ["Times New Roman", "DejaVu Serif"],
        "font.size": 9,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 300,
    }
)


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, float):
                cell.text = "" if np.isnan(value) else f"{value:.6f}"
            else:
                cell.text = str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def fit_fwl(panel: pd.DataFrame, outcome: str, treatment: np.ndarray) -> dict[str, float]:
    fixed_effects = panel[["firm_id", "industry_year_id"]].to_numpy(dtype=np.int64)
    algorithm = create(fixed_effects, drop_singletons=False)
    y = panel[outcome].to_numpy(dtype=float)
    transformed = algorithm.residualize(np.column_stack([y, treatment]))
    y_tilde, x_tilde = transformed[:, 0], transformed[:, 1]
    fit = sm.OLS(y_tilde, x_tilde[:, None], hasconst=False).fit(
        cov_type="cluster",
        cov_kwds={"groups": panel["firm_id"].to_numpy(), "use_correction": True},
        use_t=True,
    )
    return {
        "estimate": float(fit.params[0]),
        "std_error": float(fit.bse[0]),
        "t_stat": float(fit.tvalues[0]),
        "p_value": float(fit.pvalues[0]),
        "ci_low": float(fit.conf_int()[0, 0]),
        "ci_high": float(fit.conf_int()[0, 1]),
        "observations": int(len(panel)),
        "firm_clusters": int(panel["firm_id"].nunique()),
    }


def main() -> None:
    panel = pd.read_parquet(DATA / "architecture_landmark_panel.parquet").sort_values(
        ["firm_id", "year"]
    ).reset_index(drop=True)
    pretrend = json.loads((LOGS / "pretrend_gate.json").read_text(encoding="utf-8"))
    if not pretrend["gate_pass"]:
        raise RuntimeError("The frozen pretrend gate did not pass.")
    post = panel["year"].ge(2022).to_numpy(dtype=float)
    architecture = panel["architecture_count"].to_numpy(dtype=float)
    multi = panel["multi_architecture"].to_numpy(dtype=float)
    specs = [
        ("Primary: active links", "asinh_active_rcep_supplier_links", post * architecture),
        ("Binary: two-plus architectures", "asinh_active_rcep_supplier_links", post * multi),
        ("Active RCEP countries", "asinh_active_rcep_supplier_countries", post * architecture),
        ("RCEP share of Asian links", "rcep_asian_supplier_share", post * architecture),
    ]
    rows = []
    for name, outcome, treatment in specs:
        rows.append(
            {
                "specification": name,
                "outcome": outcome,
                "exposure": "multi_architecture" if "Binary" in name else "architecture_count",
                **fit_fwl(panel, outcome, treatment),
            }
        )
    primary = rows[0]

    firm = panel.drop_duplicates("firm_id").sort_values("firm_id").reset_index(drop=True)
    if not np.array_equal(firm["firm_id"].to_numpy(), np.arange(len(firm))):
        raise ValueError("Firm identifiers must be contiguous for permutation expansion.")
    rng = np.random.default_rng(SEED)
    exposure_matrix = np.empty((len(firm), B), dtype=np.int8)
    for _, stratum in firm.groupby("permutation_stratum", sort=True):
        indices = stratum.index.to_numpy()
        values = stratum["architecture_count"].to_numpy(dtype=np.int8)
        for draw in range(B):
            exposure_matrix[indices, draw] = rng.permutation(values)
    assignment_long = pd.DataFrame(
        {
            "firm_id": np.tile(firm["firm_id"].to_numpy(dtype=np.int32), B),
            "draw": np.repeat(np.arange(1, B + 1, dtype=np.int16), len(firm)),
            "permuted_architecture_count": exposure_matrix.T.reshape(-1),
        }
    )
    assignment_long.to_parquet(DATA / "stratified_permutation_assignments.parquet", index=False)

    fixed_effects = panel[["firm_id", "industry_year_id"]].to_numpy(dtype=np.int64)
    algorithm = create(fixed_effects, drop_singletons=False)
    y = panel["asinh_active_rcep_supplier_links"].to_numpy(dtype=float)
    y_tilde = algorithm.residualize(y[:, None])[:, 0]
    permutation_beta = np.empty(B, dtype=float)
    for start in range(0, B, 100):
        stop = min(start + 100, B)
        x_block = post[:, None] * exposure_matrix[panel["firm_id"].to_numpy(), start:stop]
        x_tilde = algorithm.residualize(x_block)
        permutation_beta[start:stop] = (y_tilde @ x_tilde) / np.square(x_tilde).sum(axis=0)
    permutation_p = float(
        (1 + np.sum(np.abs(permutation_beta) >= abs(primary["estimate"]))) / (B + 1)
    )
    permutation_draws = pd.DataFrame(
        {"draw": np.arange(1, B + 1), "estimate": permutation_beta}
    )
    permutation_draws.to_csv(DATA / "stratified_permutation_estimates.csv", index=False)
    permutation_draws.to_parquet(DATA / "stratified_permutation_estimates.parquet", index=False)
    primary["stratified_randomization_p"] = permutation_p

    influence_rows = []
    industries = sorted(panel["pre_industry"].unique())
    for industry in industries:
        subset = panel.loc[~panel["pre_industry"].eq(industry)].copy()
        estimate = fit_fwl(
            subset,
            "asinh_active_rcep_supplier_links",
            subset["year"].ge(2022).to_numpy(dtype=float)
            * subset["architecture_count"].to_numpy(dtype=float),
        )
        influence_rows.append(
            {
                "excluded_industry": str(industry),
                **estimate,
                "sign_reversal": bool(estimate["estimate"] * primary["estimate"] <= 0),
            }
        )
    influence = pd.DataFrame(influence_rows)
    no_sign_reversal = bool(~influence["sign_reversal"].any())

    same_binary_sign = bool(rows[1]["estimate"] * primary["estimate"] > 0)
    same_country_outcome_sign = bool(rows[2]["estimate"] * primary["estimate"] > 0)
    positive_exploratory = bool(
        permutation_p < 0.05
        and same_binary_sign
        and same_country_outcome_sign
        and no_sign_reversal
    )
    results = pd.DataFrame(rows)
    results.loc[0, "stratified_randomization_p"] = permutation_p
    export_table(results, "table2_static_results", "Cross-Architecture Static Results")
    export_table(influence, "table3_leave_one_industry_out", "Leave-One-Industry-Out Results")

    FIGURES.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.hist(permutation_beta, bins=35, color="#B8B8B8", edgecolor="white")
    ax.axvline(primary["estimate"], color="#D55E00", linewidth=1.4, label="Observed")
    ax.axvline(-primary["estimate"], color="#D55E00", linewidth=1, linestyle="--")
    ax.set_xlabel("Permuted post interaction estimate")
    ax.set_ylabel("Frequency")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig3_stratified_permutation.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig3_stratified_permutation.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    y_positions = np.arange(len(results))
    ax.errorbar(
        results["estimate"],
        y_positions,
        xerr=[results["estimate"] - results["ci_low"], results["ci_high"] - results["estimate"]],
        fmt="o",
        color="#0072B2",
        ecolor="#555555",
        capsize=3,
    )
    ax.axvline(0, color="black", linewidth=0.8)
    ax.set_yticks(y_positions, results["specification"])
    ax.set_xlabel("Post-RCEP differential estimate")
    fig.tight_layout()
    fig.savefig(FIGURES / "fig4_specification_coefficients.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig4_specification_coefficients.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    assignment_hash = hashlib.sha256(
        (DATA / "stratified_permutation_assignments.parquet").read_bytes()
    ).hexdigest()
    diagnostics = {
        "experiment_id": "09_cross_architecture_landmark",
        "gate": "STATIC_AND_STRATIFIED_RANDOMIZATION",
        "primary_result": primary,
        "permutations": {
            "draws": B,
            "seed": SEED,
            "strata": "frozen industry by size tercile",
            "two_sided_p_value": permutation_p,
            "assignment_file": "data/stratified_permutation_assignments.parquet",
            "assignment_file_sha256": assignment_hash,
        },
        "same_binary_exposure_sign": same_binary_sign,
        "same_active_country_outcome_sign": same_country_outcome_sign,
        "leave_one_industry_out_no_sign_reversal": no_sign_reversal,
        "positive_exploratory_result": positive_exploratory,
        "interpretation_limit": "Support-informed exploratory result with only one clean pre-policy lead.",
    }
    (LOGS / "static_permutation.json").write_text(json.dumps(diagnostics, indent=2), encoding="utf-8")
    status = json.loads(STATUS.read_text(encoding="utf-8"))
    status["phase"] = "STATIC_AND_STRATIFIED_RANDOMIZATION_RUN"
    status["positive_exploratory_result"] = positive_exploratory
    status["next_gate"] = "FINALIZE_VIABLE_EXPLORATORY" if positive_exploratory else "ARCHIVE_INCONCLUSIVE"
    STATUS.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(diagnostics, indent=2))


if __name__ == "__main__":
    main()
