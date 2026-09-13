#!/usr/bin/env python3
"""Build and gate the non-overlapping institutional-readiness firm panel."""

from __future__ import annotations

import hashlib
import json
import platform
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


EXPERIMENT = Path(__file__).resolve().parents[1]
REVISION = EXPERIMENT.parents[1]
WORKSPACE = REVISION.parent
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"
RCEP = {"JP", "KR", "AU", "NZ", "ID", "MY", "PH", "SG", "TH", "VN", "BN", "KH", "LA", "MM"}
ASIA_CONTROLS = {"HK", "TW", "IN", "PK", "BD", "LK", "MN"}
OUTCOME_YEARS = list(range(2020, 2025))
EXPOSURE_YEARS = [2017, 2018, 2019]


def find_legacy_root() -> Path:
    for candidate in WORKSPACE.parent.iterdir():
        path = candidate / "data" / "cleaned" / "ddd_panel.parquet"
        if path.exists():
            return candidate.resolve()
    raise FileNotFoundError("Could not locate the verified legacy data root.")


def normalize_ticker(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    values = values.where(values.str.fullmatch(r"\d{1,6}"), pd.NA)
    return values.str.zfill(6)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = f"{value:.6f}" if isinstance(value, float) else str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    pair_path = REVISION / "data" / "derived" / "pair_year.parquet"
    firm_path = find_legacy_root() / "data" / "cleaned" / "ddd_panel.parquet"
    pair = pd.read_parquet(
        pair_path,
        columns=[
            "partner_id", "partner_country", "supplier_link", "pair_id",
            "year", "active", "ticker_key",
        ],
    )
    firm = pd.read_parquet(firm_path, columns=["Stkcd", "year", "Indcat", "Size"])
    if pair.duplicated(["pair_id", "year"]).any() or not pair["active"].isin([0, 1]).all():
        raise ValueError("Relationship source failed key or active-value validation.")
    if firm.duplicated(["Stkcd", "year"]).any():
        raise ValueError("Firm source failed its firm-year key validation.")

    pair["listed_ticker"] = normalize_ticker(pair["ticker_key"])
    firm["listed_ticker"] = normalize_ticker(firm["Stkcd"])
    listed_tickers = set(firm["listed_ticker"].dropna())
    valid_pair = pair.loc[
        pair["listed_ticker"].isin(listed_tickers)
        & ~pair["listed_ticker"].eq("000000")
        & pair["partner_country"].astype("string").str.fullmatch(r"[A-Z]{2}", na=False)
    ].copy()
    covered_tickers = sorted(valid_pair["listed_ticker"].unique())

    mapping = pd.read_csv(DATA / "institutional_architecture_mapping.csv")
    if set(mapping["partner_country"]) != RCEP or mapping["partner_country"].duplicated().any():
        raise ValueError("Institutional architecture mapping must cover each RCEP country once.")
    architecture_map = mapping.set_index("partner_country")["architecture"]
    exposure_links = valid_pair.loc[
        valid_pair["supplier_link"].eq(1)
        & valid_pair["active"].eq(1)
        & valid_pair["year"].isin(EXPOSURE_YEARS)
        & valid_pair["partner_country"].isin(RCEP)
    ].copy()
    exposure_links["architecture"] = exposure_links["partner_country"].map(architecture_map)
    exposure = exposure_links.groupby("listed_ticker").agg(
        architecture_count=("architecture", "nunique"),
        rcep_country_count=("partner_country", "nunique"),
        pre_rcep_supplier_partners=("partner_id", "nunique"),
    )

    active_supplier = valid_pair.loc[
        valid_pair["supplier_link"].eq(1)
        & valid_pair["active"].eq(1)
        & valid_pair["year"].isin(OUTCOME_YEARS)
    ].copy()
    rcep_outcome = (
        active_supplier.loc[active_supplier["partner_country"].isin(RCEP)]
        .groupby(["listed_ticker", "year"], as_index=False)
        .agg(
            active_rcep_supplier_links=("partner_id", "nunique"),
            active_rcep_supplier_countries=("partner_country", "nunique"),
        )
    )
    asian_outcome = (
        active_supplier.loc[active_supplier["partner_country"].isin(RCEP | ASIA_CONTROLS)]
        .groupby(["listed_ticker", "year"], as_index=False)
        .agg(active_asian_supplier_links=("partner_id", "nunique"))
    )

    firm_panel = firm.loc[
        firm["listed_ticker"].isin(covered_tickers) & firm["year"].isin(OUTCOME_YEARS)
    ].copy()
    balanced_tickers = (
        firm_panel.groupby("listed_ticker")["year"].nunique().loc[lambda x: x.eq(5)].index
    )
    firm_panel = firm_panel.loc[firm_panel["listed_ticker"].isin(balanced_tickers)].copy()
    pre_characteristics = (
        firm.loc[
            firm["listed_ticker"].isin(balanced_tickers)
            & firm["year"].isin(EXPOSURE_YEARS)
        ]
        .sort_values(["listed_ticker", "year"])
        .groupby("listed_ticker", as_index=False)
        .agg(
            pre_industry=("Indcat", lambda x: x.dropna().iloc[-1] if x.notna().any() else np.nan),
            pre_size=("Size", "mean"),
        )
    )
    eligible_pre = pre_characteristics.dropna(subset=["pre_industry", "pre_size"])["listed_ticker"]
    firm_panel = firm_panel.loc[firm_panel["listed_ticker"].isin(eligible_pre)].copy()
    firm_panel = firm_panel.merge(pre_characteristics, on="listed_ticker", how="left", validate="many_to_one")
    firm_panel = firm_panel.merge(exposure, on="listed_ticker", how="left", validate="many_to_one")
    firm_panel = firm_panel.merge(rcep_outcome, on=["listed_ticker", "year"], how="left", validate="one_to_one")
    firm_panel = firm_panel.merge(asian_outcome, on=["listed_ticker", "year"], how="left", validate="one_to_one")
    count_columns = [
        "architecture_count", "rcep_country_count", "pre_rcep_supplier_partners",
        "active_rcep_supplier_links", "active_rcep_supplier_countries",
        "active_asian_supplier_links",
    ]
    firm_panel[count_columns] = firm_panel[count_columns].fillna(0)
    firm_panel[["architecture_count", "rcep_country_count", "pre_rcep_supplier_partners"]] = firm_panel[["architecture_count", "rcep_country_count", "pre_rcep_supplier_partners"]].astype("int16")
    firm_panel[["active_rcep_supplier_links", "active_rcep_supplier_countries", "active_asian_supplier_links"]] = firm_panel[["active_rcep_supplier_links", "active_rcep_supplier_countries", "active_asian_supplier_links"]].astype("int32")
    firm_panel["multi_architecture"] = firm_panel["architecture_count"].ge(2).astype("int8")
    firm_panel["asinh_active_rcep_supplier_links"] = np.arcsinh(firm_panel["active_rcep_supplier_links"])
    firm_panel["asinh_active_rcep_supplier_countries"] = np.arcsinh(firm_panel["active_rcep_supplier_countries"])
    firm_panel["rcep_asian_supplier_share"] = np.divide(
        firm_panel["active_rcep_supplier_links"],
        firm_panel["active_asian_supplier_links"],
        out=np.zeros(len(firm_panel), dtype=float),
        where=firm_panel["active_asian_supplier_links"].gt(0),
    )
    firm_panel["post2022"] = firm_panel["year"].ge(2022).astype("int8")
    firm_panel["post_architecture"] = firm_panel["post2022"] * firm_panel["architecture_count"]
    firm_panel["post_multi_architecture"] = firm_panel["post2022"] * firm_panel["multi_architecture"]

    firm_level = firm_panel.drop_duplicates("listed_ticker").copy()
    firm_level["size_tercile"] = pd.qcut(
        firm_level["pre_size"], q=3, labels=[0, 1, 2], duplicates="drop"
    ).astype("int8")
    firm_level["permutation_stratum"] = (
        firm_level["pre_industry"].astype(str) + "|" + firm_level["size_tercile"].astype(str)
    )
    firm_panel = firm_panel.drop(columns=["size_tercile", "permutation_stratum"], errors="ignore").merge(
        firm_level[["listed_ticker", "size_tercile", "permutation_stratum"]],
        on="listed_ticker",
        how="left",
        validate="many_to_one",
    )
    firm_panel["firm_id"] = pd.factorize(firm_panel["listed_ticker"], sort=True)[0].astype("int64")
    firm_panel["industry_year_id"] = pd.factorize(
        pd.MultiIndex.from_frame(firm_panel[["pre_industry", "year"]]), sort=True
    )[0].astype("int64")

    key = ["listed_ticker", "year"]
    balanced = firm_panel.groupby("listed_ticker")["year"].nunique().eq(5).all()
    if firm_panel.duplicated(key).any() or not balanced:
        raise ValueError("Final five-year firm panel is not unique and balanced.")
    firm_level = firm_panel.drop_duplicates("listed_ticker")
    exposed = int(firm_level["architecture_count"].ge(1).sum())
    multi = int(firm_level["multi_architecture"].sum())
    panel_firms = int(firm_level["listed_ticker"].nunique())
    positive_by_year = firm_panel.groupby("year")["active_rcep_supplier_links"].apply(lambda x: int(x.gt(0).sum()))
    industry_count = int(firm_level["pre_industry"].nunique())
    stratum_sizes = firm_level["permutation_stratum"].value_counts()
    singleton_strata = set(stratum_sizes.loc[stratum_sizes.eq(1)].index)
    exposed_singletons = int(
        firm_level.loc[
            firm_level["permutation_stratum"].isin(singleton_strata)
            & firm_level["architecture_count"].gt(0)
        ].shape[0]
    )
    exposure_gate = exposed >= 250 and multi >= 40
    panel_gate = panel_firms >= 3000 and bool(positive_by_year.ge(100).all()) and balanced
    strata_gate = industry_count >= 15 and exposed_singletons <= 10
    gate_pass = bool(exposure_gate and panel_gate and strata_gate)

    exposure_distribution = firm_level.groupby("architecture_count", as_index=False).agg(
        firms=("listed_ticker", "nunique"),
        mean_pre_rcep_countries=("rcep_country_count", "mean"),
        mean_pre_supplier_partners=("pre_rcep_supplier_partners", "mean"),
    )
    year_support = firm_panel.groupby("year", as_index=False).agg(
        firms=("listed_ticker", "nunique"),
        positive_primary=("active_rcep_supplier_links", lambda x: int(x.gt(0).sum())),
        active_rcep_supplier_links=("active_rcep_supplier_links", "sum"),
    )
    support_table = pd.DataFrame(
        [
            {"metric": "Firms with >=1 architecture", "value": exposed, "threshold": 250, "pass": exposed >= 250},
            {"metric": "Firms with >=2 architectures", "value": multi, "threshold": 40, "pass": multi >= 40},
            {"metric": "Balanced panel firms", "value": panel_firms, "threshold": 3000, "pass": panel_firms >= 3000},
            {"metric": "Minimum annual positive outcomes", "value": int(positive_by_year.min()), "threshold": 100, "pass": bool(positive_by_year.ge(100).all())},
            {"metric": "Pre-policy industries", "value": industry_count, "threshold": 15, "pass": industry_count >= 15},
            {"metric": "Exposed firms in singleton strata", "value": exposed_singletons, "threshold": 10, "pass": exposed_singletons <= 10},
        ]
    )

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    firm_panel = firm_panel.sort_values(["firm_id", "year"]).reset_index(drop=True)
    firm_panel.to_parquet(DATA / "architecture_firm_year.parquet", index=False)
    exposure_distribution.to_csv(DATA / "exposure_distribution.csv", index=False)
    year_support.to_csv(DATA / "year_outcome_support.csv", index=False)
    export_table(support_table, "table1_support_gate", "Institutional Readiness Support Gate")
    gate = {
        "experiment_id": "07_fta_novelty",
        "gate": "DATA_SUPPORT",
        "gate_pass": gate_pass,
        "exposure_gate": exposure_gate,
        "panel_gate": panel_gate,
        "strata_gate": strata_gate,
        "firms_with_one_or_more_architectures": exposed,
        "firms_with_two_or_more_architectures": multi,
        "balanced_panel_firms": panel_firms,
        "minimum_annual_positive_primary_outcome": int(positive_by_year.min()),
        "pre_policy_industries": industry_count,
        "permutation_strata": int(len(stratum_sizes)),
        "exposed_firms_in_singleton_strata": exposed_singletons,
    }
    (DATA / "support_gate.json").write_text(json.dumps(gate, indent=2), encoding="utf-8")
    summary = {
        "experiment_id": "07_fta_novelty",
        "build_status": "PASSED",
        "python": platform.python_version(),
        "source_hashes": {
            "pair_year_sha256": sha256(pair_path),
            "firm_year_sha256": sha256(firm_path),
        },
        "valid_relationship_rows": int(len(valid_pair)),
        "covered_tickers": int(len(covered_tickers)),
        "support_gate": gate,
    }
    (DATA / "panel_build_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    (LOGS / "panel_build.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    sample_path = DATA / "sample_construction.json"
    sample = json.loads(sample_path.read_text(encoding="utf-8"))
    sample["status"] = "COMPLETED"
    sample["counts"] = summary
    sample["execution_log"] = "logs/panel_build.json"
    sample_path.write_text(json.dumps(sample, indent=2), encoding="utf-8")
    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "DATA_SUPPORT_GATE_RUN"
    status["results_inspected"] = True
    status["support_gate_pass"] = gate_pass
    status["next_gate"] = "PRETREND" if gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
