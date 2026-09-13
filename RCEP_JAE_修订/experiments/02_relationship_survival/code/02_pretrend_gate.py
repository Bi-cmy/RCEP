#!/usr/bin/env python3
"""Run the frozen event-study pre-trend gate for Experiment 02."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from linearmodels.panel import PanelOLS
from scipy import stats


EXPERIMENT = Path(__file__).resolve().parents[1]
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
FIGURES = EXPERIMENT / "figures"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"


plt.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": ["Times New Roman", "DejaVu Serif"],
        "font.size": 9,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 300,
    }
)


def term_name(group: int, event_time: int) -> str:
    label = f"m{abs(event_time)}" if event_time < 0 else f"p{event_time}"
    return f"g{group}_e_{label}"


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, float):
                cell.text = "" if np.isnan(value) else f"{value:.6f}"
            else:
                cell.text = str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    panel = pd.read_parquet(DATA / "survival_cohort.parquet")
    if panel.duplicated(["relationship_id", "year"]).any():
        raise ValueError("Survival cohort has duplicate relationship-year keys.")

    support = pd.read_csv(DATA / "group_year_support.csv")
    table1 = support.rename(
        columns={
            "rcep": "RCEP cohort",
            "year": "Year",
            "countries": "Countries",
            "relationships": "Relationships",
            "listed_firms": "Listed firms",
            "survivors": "Survivors",
            "at_risk": "At risk",
            "first_exits": "First exits",
        }
    )
    export_table(table1, "table1_cohort_support", "Fixed 2017 Cohort Support")

    work = panel.copy()
    cohort_sizes = (
        work.loc[work["rcep"].eq(1), ["relationship_id", "entry_year"]]
        .drop_duplicates()
        .groupby("entry_year")["relationship_id"]
        .nunique()
        .to_dict()
    )
    terms: list[str] = []
    term_meta: list[dict[str, int]] = []
    for group in sorted(cohort_sizes):
        event_times = sorted(
            work.loc[work["entry_year"].eq(group), "event_time"].dropna().astype(int).unique()
        )
        for event_time in event_times:
            if event_time == -1:
                continue
            name = term_name(group, event_time)
            work[name] = (
                work["entry_year"].eq(group) & work["event_time"].eq(event_time)
            ).astype("int8")
            terms.append(name)
            term_meta.append(
                {"term": name, "group": int(group), "event_time": int(event_time)}
            )

    indexed = work.sort_values(["relationship_id", "year"]).set_index(
        ["relationship_id", "year"]
    )
    clusters = pd.DataFrame(
        {"partner_country": pd.Categorical(indexed["partner_country"]).codes},
        index=indexed.index,
    )
    model = PanelOLS(
        indexed["survival"].astype(float),
        indexed[terms].astype(float),
        entity_effects=True,
        time_effects=True,
        drop_absorbed=False,
        check_rank=True,
    )
    result = model.fit(cov_type="clustered", clusters=clusters)
    missing_terms = sorted(set(terms).difference(result.params.index))
    if missing_terms:
        raise ValueError(f"Event-study terms were dropped: {missing_terms}")

    pre_terms = [meta["term"] for meta in term_meta if meta["event_time"] <= -2]
    restriction = np.zeros((len(pre_terms), len(result.params)))
    parameter_names = list(result.params.index)
    for row, term in enumerate(pre_terms):
        restriction[row, parameter_names.index(term)] = 1
    wald = result.wald_test(restriction)
    pretrend_p = float(np.asarray(wald.pval).squeeze())
    pretrend_stat = float(np.asarray(wald.stat).squeeze())

    rows = []
    cluster_count = int(work["partner_country"].nunique())
    critical = float(stats.t.ppf(0.975, df=cluster_count - 1))
    for event_time in sorted({meta["event_time"] for meta in term_meta}):
        meta_at_time = [meta for meta in term_meta if meta["event_time"] == event_time]
        names = [meta["term"] for meta in meta_at_time]
        raw_weights = np.array([cohort_sizes[meta["group"]] for meta in meta_at_time], dtype=float)
        weights = raw_weights / raw_weights.sum()
        beta = result.params.loc[names].to_numpy(dtype=float)
        covariance = result.cov.loc[names, names].to_numpy(dtype=float)
        estimate = float(weights @ beta)
        se = float(np.sqrt(max(weights @ covariance @ weights, 0.0)))
        t_stat = estimate / se if se > 0 else np.nan
        p_value = float(2 * stats.t.sf(abs(t_stat), df=cluster_count - 1)) if se > 0 else np.nan
        rows.append(
            {
                "event_time": int(event_time),
                "estimate": estimate,
                "std_error": se,
                "ci_low": estimate - critical * se,
                "ci_high": estimate + critical * se,
                "p_value_t14": p_value,
                "contributing_cohorts": ",".join(str(meta["group"]) for meta in meta_at_time),
            }
        )
    rows.append(
        {
            "event_time": -1,
            "estimate": 0.0,
            "std_error": 0.0,
            "ci_low": 0.0,
            "ci_high": 0.0,
            "p_value_t14": np.nan,
            "contributing_cohorts": "reference",
        }
    )
    event = pd.DataFrame(rows).sort_values("event_time").reset_index(drop=True)
    pre_event = event.loc[event["event_time"].le(-2)]
    max_abs_pre = float(pre_event["estimate"].abs().max())
    gate_pass = bool(pretrend_p >= 0.10 and max_abs_pre <= 0.10)

    export_table(event, "table2_event_study_gate", "Event-Study Identification Gate")
    event.to_csv(DATA / "event_study_gate.csv", index=False)

    FIGURES.mkdir(parents=True, exist_ok=True)
    trends = (
        work.groupby(["year", "rcep"], as_index=False)["survival"].mean()
        .pivot(index="year", columns="rcep", values="survival")
    )
    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.plot(trends.index, trends[0], marker="s", linestyle="--", color="#555555", label="Asian non-RCEP")
    ax.plot(trends.index, trends[1], marker="o", color="#0072B2", label="RCEP")
    ax.axvline(2021.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Year")
    ax.set_ylabel("Share of fixed 2017 cohort surviving")
    ax.set_ylim(0, 1)
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig1_survival_trends.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig1_survival_trends.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    plotted = event.loc[event["event_time"].ne(-1)]
    ax.errorbar(
        plotted["event_time"],
        plotted["estimate"],
        yerr=[
            plotted["estimate"] - plotted["ci_low"],
            plotted["ci_high"] - plotted["estimate"],
        ],
        fmt="o-",
        color="#0072B2",
        ecolor="#555555",
        elinewidth=1,
        capsize=3,
        markersize=4,
        linewidth=1,
    )
    ax.scatter([-1], [0], facecolors="white", edgecolors="#0072B2", s=28, zorder=3)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.axvline(-0.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Years relative to RCEP entry")
    ax.set_ylabel("Difference in survival probability")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig2_event_study_gate.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig2_event_study_gate.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    diagnostics = {
        "experiment_id": "02_relationship_survival",
        "gate": "PRETREND",
        "gate_pass": gate_pass,
        "joint_pretrend_stat": pretrend_stat,
        "joint_pretrend_df": len(pre_terms),
        "joint_pretrend_p_value": pretrend_p,
        "maximum_absolute_aggregated_pre_coefficient": max_abs_pre,
        "economic_threshold": 0.10,
        "relationship_count": int(work["relationship_id"].nunique()),
        "partner_country_clusters": cluster_count,
        "cluster": "partner_country",
        "estimator": "Cohort-by-event interactions with relationship and year fixed effects",
        "inference_note": "Intervals use t critical values with G-1 degrees of freedom; wild bootstrap is gated on pretrend passage.",
    }
    (LOGS / "pretrend_gate.json").write_text(
        json.dumps(diagnostics, indent=2), encoding="utf-8"
    )
    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "PRETREND_GATE_RUN"
    status["results_inspected"] = True
    status["pretrend_gate_pass"] = gate_pass
    status["next_gate"] = "WILD_BOOTSTRAP_AND_RANDOMIZATION" if gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(diagnostics, indent=2))


if __name__ == "__main__":
    main()
