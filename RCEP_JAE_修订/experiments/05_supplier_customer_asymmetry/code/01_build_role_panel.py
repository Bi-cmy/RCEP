#!/usr/bin/env python3
"""Build and gate the balanced supplier-customer role panel."""

from __future__ import annotations

import hashlib
import json
import platform
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


EXPERIMENT = Path(__file__).resolve().parents[1]
REVISION = EXPERIMENT.parents[1]
WORKSPACE = REVISION.parent
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"
RCEP = {"JP", "KR", "AU", "NZ", "ID", "MY", "PH", "SG", "TH", "VN", "BN", "KH", "LA", "MM"}
LATE_2023 = {"ID", "PH"}
ASIA_CONTROLS = {"HK", "TW", "IN", "PK", "BD", "LK", "MN"}
COUNTRIES = sorted(RCEP | ASIA_CONTROLS)
ROLES = [0, 1]
YEARS = list(range(2017, 2025))
ANALYSIS_YEARS = list(range(2018, 2025))


def find_legacy_root() -> Path:
    for candidate in WORKSPACE.parent.iterdir():
        path = candidate / "data" / "cleaned" / "ddd_panel.parquet"
        if path.exists():
            return candidate.resolve()
    raise FileNotFoundError("Could not locate the verified legacy data root.")


def normalize_ticker(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
    values = values.where(values.str.fullmatch(r"\d{1,6}"), pd.NA)
    return values.str.zfill(6)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = f"{value:.6f}" if isinstance(value, float) else str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    pair_path = REVISION / "data" / "derived" / "pair_year.parquet"
    firm_path = find_legacy_root() / "data" / "cleaned" / "ddd_panel.parquet"
    pair = pd.read_parquet(
        pair_path,
        columns=[
            "cn_id", "partner_id", "partner_country", "supplier_link", "pair_id",
            "year", "active", "ticker_key",
        ],
    )
    firm = pd.read_parquet(firm_path, columns=["Stkcd", "year"])
    if pair.duplicated(["pair_id", "year"]).any() or not pair["active"].isin([0, 1]).all():
        raise ValueError("Pair-year source failed key or active-value validation.")
    if not pair["supplier_link"].isin([0, 1]).all():
        raise ValueError("Relationship role is not binary.")

    pair["listed_ticker"] = normalize_ticker(pair["ticker_key"])
    firm["listed_ticker"] = normalize_ticker(firm["Stkcd"])
    listed_tickers = set(firm["listed_ticker"].dropna())
    mapped = pair.loc[
        pair["listed_ticker"].isin(listed_tickers)
        & ~pair["listed_ticker"].eq("000000")
    ].copy()
    covered_firms = sorted(
        mapped.loc[
            mapped["year"].eq(2017)
            & mapped["active"].eq(1)
            & mapped["partner_country"].astype("string").str.fullmatch(r"[A-Z]{2}", na=False),
            "listed_ticker",
        ].unique()
    )
    if not covered_firms:
        raise ValueError("No firms satisfy the frozen 2017 relationship-coverage rule.")

    unique_partner = (
        mapped.loc[
            mapped["listed_ticker"].isin(covered_firms)
            & mapped["partner_country"].isin(COUNTRIES)
        ]
        .groupby(
            [
                "listed_ticker", "partner_country", "supplier_link", "partner_id", "year"
            ],
            as_index=False,
        )["active"]
        .max()
    )
    counts = (
        unique_partner.groupby(
            ["listed_ticker", "partner_country", "supplier_link", "year"], as_index=False
        )["active"]
        .sum()
        .rename(
            columns={"supplier_link": "supplier_role", "active": "active_links"}
        )
    )
    panel = pd.MultiIndex.from_product(
        [covered_firms, COUNTRIES, ROLES, ANALYSIS_YEARS],
        names=["listed_ticker", "partner_country", "supplier_role", "year"],
    ).to_frame(index=False)
    panel = panel.merge(
        counts,
        on=["listed_ticker", "partner_country", "supplier_role", "year"],
        how="left",
        validate="one_to_one",
    )
    panel["active_links"] = panel["active_links"].fillna(0).astype("int32")
    panel["any_active_link"] = panel["active_links"].gt(0).astype("int8")
    panel["asinh_active_links"] = np.arcsinh(panel["active_links"])
    panel["rcep"] = panel["partner_country"].isin(RCEP).astype("int8")
    panel["entry_policy_year"] = np.where(
        panel["partner_country"].isin(LATE_2023),
        2023,
        np.where(panel["rcep"].eq(1), 2022, 9999),
    ).astype("int16")
    panel["treated_staged"] = (
        panel["rcep"].eq(1) & panel["year"].ge(panel["entry_policy_year"])
    ).astype("int8")
    panel["treated_uniform"] = (
        panel["rcep"].eq(1) & panel["year"].ge(2022)
    ).astype("int8")
    panel["triple_staged"] = (panel["supplier_role"] * panel["treated_staged"]).astype("int8")
    panel["triple_uniform"] = (panel["supplier_role"] * panel["treated_uniform"]).astype("int8")
    panel["event_time"] = np.where(
        panel["rcep"].eq(1), panel["year"] - panel["entry_policy_year"], np.nan
    )
    panel["period"] = np.where(panel["year"].le(2021), "pre", "post")
    panel["firm_country_role_id"] = pd.factorize(
        pd.MultiIndex.from_frame(
            panel[["listed_ticker", "partner_country", "supplier_role"]]
        ),
        sort=True,
    )[0].astype("int64")
    panel["firm_year_role_id"] = pd.factorize(
        pd.MultiIndex.from_frame(panel[["listed_ticker", "year", "supplier_role"]]),
        sort=True,
    )[0].astype("int64")
    panel["country_year_id"] = pd.factorize(
        pd.MultiIndex.from_frame(panel[["partner_country", "year"]]), sort=True
    )[0].astype("int64")

    key = ["listed_ticker", "partner_country", "supplier_role", "year"]
    if panel.duplicated(key).any():
        raise ValueError("Role-panel key is not unique.")
    observations_per_cell = panel.groupby(key[:-1])["year"].nunique()
    if not observations_per_cell.eq(len(ANALYSIS_YEARS)).all():
        raise ValueError("Balanced role-panel validation failed.")

    support = (
        panel.groupby(["supplier_role", "rcep", "period"], as_index=False)
        .agg(
            rows=("active_links", "size"),
            firms=("listed_ticker", "nunique"),
            countries=("partner_country", "nunique"),
            active_links=("active_links", "sum"),
            positive_cells=("any_active_link", "sum"),
        )
    )
    country_support = (
        panel.groupby(
            ["supplier_role", "partner_country", "rcep", "period"], as_index=False
        )
        .agg(active_links=("active_links", "sum"))
    )
    firm_period_role = (
        panel.groupby(["listed_ticker", "period", "supplier_role"], as_index=False)[
            "active_links"
        ]
        .sum()
        .pivot_table(
            index=["listed_ticker", "period"],
            columns="supplier_role",
            values="active_links",
            fill_value=0,
        )
        .reset_index()
    )
    for column in ROLES:
        if column not in firm_period_role:
            firm_period_role[column] = 0
    firm_period_role["positive_both_roles"] = (
        firm_period_role[0].gt(0) & firm_period_role[1].gt(0)
    )
    firms_both_roles = (
        firm_period_role.groupby("period", as_index=False)["positive_both_roles"]
        .sum()
        .rename(columns={"positive_both_roles": "firms_positive_in_both_roles"})
    )

    support_key = support.set_index(["supplier_role", "rcep", "period"])["active_links"]
    required_keys = [
        (role, rcep, period)
        for role in ROLES
        for rcep in [0, 1]
        for period in ["pre", "post"]
    ]
    link_gate = all(int(support_key.get(key, 0)) >= 100 for key in required_keys)
    positive_countries = country_support.loc[country_support["active_links"].gt(0)].groupby(
        ["supplier_role", "rcep", "period"]
    )["partner_country"].nunique()
    country_gate = all(int(positive_countries.get(key, 0)) >= 5 for key in required_keys)
    firms_key = firms_both_roles.set_index("period")["firms_positive_in_both_roles"]
    firm_gate = all(int(firms_key.get(period, 0)) >= 50 for period in ["pre", "post"])
    gate_pass = bool(link_gate and country_gate and firm_gate)

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    panel = panel.sort_values(["firm_country_role_id", "year"]).reset_index(drop=True)
    panel.to_parquet(DATA / "firm_country_role_year.parquet", index=False)
    support.to_csv(DATA / "support_by_role_group_period.csv", index=False)
    country_support.to_csv(DATA / "country_role_period_support.csv", index=False)
    firms_both_roles.to_csv(DATA / "firm_role_overlap.csv", index=False)
    export_table(support, "table1_support_gate", "Supplier-Customer Role Support")

    gate = {
        "experiment_id": "05_supplier_customer_asymmetry",
        "gate": "DATA_SUPPORT",
        "gate_pass": gate_pass,
        "active_link_gate": link_gate,
        "positive_country_gate": country_gate,
        "positive_firm_both_roles_gate": firm_gate,
        "active_links": {
            f"role_{key[0]}_rcep_{key[1]}_{key[2]}": int(support_key.get(key, 0))
            for key in required_keys
        },
        "positive_countries": {
            f"role_{key[0]}_rcep_{key[1]}_{key[2]}": int(positive_countries.get(key, 0))
            for key in required_keys
        },
        "firms_positive_in_both_roles": {
            period: int(firms_key.get(period, 0)) for period in ["pre", "post"]
        },
    }
    (DATA / "support_gate.json").write_text(json.dumps(gate, indent=2), encoding="utf-8")
    summary = {
        "experiment_id": "05_supplier_customer_asymmetry",
        "build_status": "PASSED",
        "python": platform.python_version(),
        "source_hashes": {
            "pair_year_sha256": sha256(pair_path),
            "firm_year_sha256": sha256(firm_path),
        },
        "frozen_2017_covered_firms": int(len(covered_firms)),
        "panel_rows": int(len(panel)),
        "firm_country_role_cells": int(panel["firm_country_role_id"].nunique()),
        "support_gate": gate,
    }
    (DATA / "panel_build_summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    (LOGS / "panel_build.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    sample_path = DATA / "sample_construction.json"
    sample = json.loads(sample_path.read_text(encoding="utf-8"))
    sample["status"] = "COMPLETED"
    sample["counts"] = summary
    sample["execution_log"] = "logs/panel_build.json"
    sample_path.write_text(json.dumps(sample, indent=2), encoding="utf-8")

    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "DATA_SUPPORT_GATE_RUN"
    status["results_inspected"] = True
    status["support_gate_pass"] = gate_pass
    status["next_gate"] = "TRIPLE_DIFFERENCE_PRETREND" if gate_pass else "ARCHIVE_FAILURE"
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
