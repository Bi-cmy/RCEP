#!/usr/bin/env python3
"""Run the frozen supplier-customer triple-difference pretrend gate."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from pyhdfe import create
from scipy import stats


EXPERIMENT = Path(__file__).resolve().parents[1]
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
FIGURES = EXPERIMENT / "figures"
LOGS = EXPERIMENT / "logs"
STATUS_PATH = EXPERIMENT / "STATUS.json"


plt.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": ["Times New Roman", "DejaVu Serif"],
        "font.size": 9,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 300,
    }
)


def term_name(group: int, event_time: int) -> str:
    suffix = f"m{abs(event_time)}" if event_time < 0 else f"p{event_time}"
    return f"supplier_g{group}_e_{suffix}"


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, float):
                cell.text = "" if np.isnan(value) else f"{value:.6f}"
            else:
                cell.text = str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def main() -> None:
    panel = pd.read_parquet(DATA / "firm_country_role_year.parquet")
    key = ["listed_ticker", "partner_country", "supplier_role", "year"]
    if panel.duplicated(key).any():
        raise ValueError("Firm-country-role-year panel has duplicate keys.")
    if panel.groupby(key[:-1])["year"].nunique().ne(7).any():
        raise ValueError("The frozen seven-year balanced-panel condition failed.")

    work = panel.copy()
    treated_supplier = work.loc[
        work["rcep"].eq(1) & work["supplier_role"].eq(1),
        ["firm_country_role_id", "entry_policy_year"],
    ].drop_duplicates()
    cohort_sizes = (
        treated_supplier.groupby("entry_policy_year")["firm_country_role_id"]
        .nunique()
        .to_dict()
    )
    terms: list[str] = []
    meta: list[dict[str, int]] = []
    for group in sorted(cohort_sizes):
        event_times = sorted(
            work.loc[work["entry_policy_year"].eq(group), "event_time"]
            .dropna()
            .astype(int)
            .unique()
        )
        for event_time in event_times:
            if event_time == -1:
                continue
            term = term_name(int(group), int(event_time))
            work[term] = (
                work["supplier_role"].eq(1)
                & work["entry_policy_year"].eq(group)
                & work["event_time"].eq(event_time)
            ).astype("int8")
            terms.append(term)
            meta.append(
                {"term": term, "group": int(group), "event_time": int(event_time)}
            )

    y = work["asinh_active_links"].to_numpy(dtype=float)
    x = work[terms].to_numpy(dtype=float)
    fixed_effects = work[
        ["firm_country_role_id", "firm_year_role_id", "country_year_id"]
    ].to_numpy(dtype=np.int64)
    algorithm = create(fixed_effects, drop_singletons=False)
    residualized = algorithm.residualize(np.column_stack([y, x]))
    y_tilde = residualized[:, 0]
    x_tilde = residualized[:, 1:]
    clusters = pd.Categorical(work["partner_country"]).codes
    cluster_count = int(work["partner_country"].nunique())
    result = sm.OLS(y_tilde, x_tilde, hasconst=False).fit(
        cov_type="cluster",
        cov_kwds={"groups": clusters, "use_correction": True},
        use_t=True,
    )
    beta = pd.Series(result.params, index=terms)
    covariance = pd.DataFrame(result.cov_params(), index=terms, columns=terms)

    pre_terms = [item["term"] for item in meta if item["event_time"] <= -2]
    pre_beta = beta.loc[pre_terms].to_numpy(dtype=float)
    pre_cov = covariance.loc[pre_terms, pre_terms].to_numpy(dtype=float)
    covariance_rank = int(np.linalg.matrix_rank(pre_cov))
    wald_stat = float(pre_beta @ np.linalg.pinv(pre_cov) @ pre_beta)
    pretrend_p = float(stats.chi2.sf(wald_stat, df=covariance_rank))

    critical = float(stats.t.ppf(0.975, df=cluster_count - 1))
    rows = []
    for event_time in sorted({item["event_time"] for item in meta}):
        selected = [item for item in meta if item["event_time"] == event_time]
        names = [item["term"] for item in selected]
        raw_weights = np.array(
            [cohort_sizes[item["group"]] for item in selected], dtype=float
        )
        weights = raw_weights / raw_weights.sum()
        estimate = float(weights @ beta.loc[names].to_numpy(dtype=float))
        selected_cov = covariance.loc[names, names].to_numpy(dtype=float)
        se = float(np.sqrt(max(weights @ selected_cov @ weights, 0.0)))
        t_stat = estimate / se if se > 0 else np.nan
        p_value = (
            float(2 * stats.t.sf(abs(t_stat), df=cluster_count - 1))
            if se > 0
            else np.nan
        )
        rows.append(
            {
                "event_time": int(event_time),
                "estimate": estimate,
                "std_error": se,
                "ci_low": estimate - critical * se,
                "ci_high": estimate + critical * se,
                "p_value_t20": p_value,
                "contributing_cohorts": ",".join(
                    str(item["group"]) for item in selected
                ),
            }
        )
    rows.append(
        {
            "event_time": -1,
            "estimate": 0.0,
            "std_error": 0.0,
            "ci_low": 0.0,
            "ci_high": 0.0,
            "p_value_t20": np.nan,
            "contributing_cohorts": "reference",
        }
    )
    event = pd.DataFrame(rows).sort_values("event_time").reset_index(drop=True)
    maximum_absolute_pre = float(
        event.loc[event["event_time"].le(-2), "estimate"].abs().max()
    )
    gate_pass = bool(pretrend_p >= 0.10 and maximum_absolute_pre <= 0.05)

    DATA.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    event.to_csv(DATA / "event_study_gate.csv", index=False)
    cohort_terms = pd.DataFrame(meta)
    cohort_terms["estimate"] = cohort_terms["term"].map(beta)
    cohort_terms["std_error"] = cohort_terms["term"].map(
        pd.Series(np.sqrt(np.diag(covariance)), index=terms)
    )
    cohort_terms.to_csv(DATA / "cohort_event_terms.csv", index=False)
    export_table(
        event,
        "table2_event_study_gate",
        "Supplier-Customer Asymmetry Event Study",
    )

    FIGURES.mkdir(parents=True, exist_ok=True)
    role_means = (
        work.groupby(["year", "rcep", "supplier_role"], as_index=False)[
            "asinh_active_links"
        ]
        .mean()
        .pivot(index=["year", "rcep"], columns="supplier_role", values="asinh_active_links")
        .reset_index()
    )
    role_means["supplier_customer_gap"] = role_means[1] - role_means[0]
    trends = role_means.pivot(
        index="year", columns="rcep", values="supplier_customer_gap"
    )
    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.plot(
        trends.index,
        trends[0],
        marker="s",
        linestyle="--",
        color="#555555",
        label="Asian non-RCEP",
    )
    ax.plot(
        trends.index,
        trends[1],
        marker="o",
        color="#0072B2",
        label="RCEP",
    )
    ax.axvline(2021.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Year")
    ax.set_ylabel("Mean supplier-customer link gap")
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig1_raw_role_gap_trends.pdf", bbox_inches="tight")
    fig.savefig(
        FIGURES / "fig1_raw_role_gap_trends.png", dpi=300, bbox_inches="tight"
    )
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    plotted = event.loc[event["event_time"].ne(-1)]
    ax.errorbar(
        plotted["event_time"],
        plotted["estimate"],
        yerr=[
            plotted["estimate"] - plotted["ci_low"],
            plotted["ci_high"] - plotted["estimate"],
        ],
        fmt="o-",
        color="#0072B2",
        ecolor="#555555",
        elinewidth=1,
        capsize=3,
        markersize=4,
        linewidth=1,
    )
    ax.scatter([-1], [0], facecolors="white", edgecolors="#0072B2", s=28, zorder=3)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.axvline(-0.5, color="#D55E00", linestyle="--", linewidth=1)
    ax.set_xlabel("Years relative to RCEP entry")
    ax.set_ylabel("Supplier-customer differential response")
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig2_event_study_gate.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig2_event_study_gate.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    diagnostics = {
        "experiment_id": "05_supplier_customer_asymmetry",
        "gate": "TRIPLE_DIFFERENCE_PRETREND",
        "gate_pass": gate_pass,
        "joint_pretrend_wald_stat": wald_stat,
        "joint_pretrend_df": covariance_rank,
        "joint_pretrend_p_value": pretrend_p,
        "maximum_absolute_aggregated_pre_coefficient": maximum_absolute_pre,
        "economic_threshold": 0.05,
        "observations": int(len(work)),
        "firm_country_role_fixed_effects": int(
            work["firm_country_role_id"].nunique()
        ),
        "firm_year_role_fixed_effects": int(work["firm_year_role_id"].nunique()),
        "country_year_fixed_effects": int(work["country_year_id"].nunique()),
        "partner_country_clusters": cluster_count,
        "cohort_sizes_supplier_cells": {
            str(key): int(value) for key, value in cohort_sizes.items()
        },
        "estimator": (
            "OLS after pyhdfe residualization on firm-country-role, "
            "firm-year-role, and country-year fixed effects"
        ),
    }
    (LOGS / "pretrend_gate.json").write_text(
        json.dumps(diagnostics, indent=2), encoding="utf-8"
    )
    status = json.loads(STATUS_PATH.read_text(encoding="utf-8"))
    status["phase"] = "PRETREND_GATE_RUN"
    status["results_inspected"] = True
    status["pretrend_gate_pass"] = gate_pass
    status["next_gate"] = (
        "STATIC_EFFECT_AND_RANDOMIZATION" if gate_pass else "ARCHIVE_FAILURE"
    )
    STATUS_PATH.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(diagnostics, indent=2))


if __name__ == "__main__":
    main()
