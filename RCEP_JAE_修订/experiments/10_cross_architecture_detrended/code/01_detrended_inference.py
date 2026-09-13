#!/usr/bin/env python3
"""Test the cross-architecture result after a frozen linear differential trend."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from pyhdfe import create


EXPERIMENT = Path(__file__).resolve().parents[1]
SOURCE_EXPERIMENT = EXPERIMENT.parent / "09_cross_architecture_landmark"
SOURCE = SOURCE_EXPERIMENT / "data" / "architecture_landmark_panel.parquet"
DATA = EXPERIMENT / "data"
TABLES = EXPERIMENT / "tables"
FIGURES = EXPERIMENT / "figures"
LOGS = EXPERIMENT / "logs"
STATUS = EXPERIMENT / "STATUS.json"
SEED = 20260723
B = 999


plt.rcParams.update(
    {
        "font.family": "serif",
        "font.serif": ["Times New Roman", "DejaVu Serif"],
        "font.size": 9,
        "axes.labelsize": 10,
        "xtick.labelsize": 8,
        "ytick.labelsize": 8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.dpi": 300,
    }
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_docx(frame: pd.DataFrame, path: Path, title: str) -> None:
    document = Document()
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(frame.columns))
    for cell, column in zip(table.rows[0].cells, frame.columns, strict=True):
        cell.text = str(column)
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            if isinstance(value, float):
                cell.text = "" if np.isnan(value) else f"{value:.6f}"
            else:
                cell.text = str(value)
    document.save(path)


def export_table(frame: pd.DataFrame, stem: str, title: str) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    frame.to_excel(TABLES / f"{stem}.xlsx", index=False)
    frame.to_latex(
        TABLES / f"{stem}.tex",
        index=False,
        float_format="%.6f",
        caption=title,
        label=f"tab:{stem}",
    )
    write_docx(frame, TABLES / f"{stem}.docx", title)


def solve_columns(y: np.ndarray, x: np.ndarray) -> np.ndarray:
    xx = x.T @ x
    xy = x.T @ y
    return np.linalg.pinv(xx) @ xy


def main() -> None:
    panel = pd.read_parquet(SOURCE).sort_values(["firm_id", "year"]).reset_index(drop=True)
    if len(panel) != 12780 or panel["firm_id"].nunique() != 2556:
        raise ValueError("Experiment 09 panel counts changed.")
    DATA.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    cloned = DATA / "detrended_panel.parquet"
    panel.to_parquet(cloned, index=False)

    time = (panel["year"] - 2021).to_numpy(dtype=float)
    post = panel["year"].ge(2022).to_numpy(dtype=float)
    exposure = panel["architecture_count"].to_numpy(dtype=float)
    x = np.column_stack([exposure * time, exposure * post])
    y = panel["asinh_active_rcep_supplier_links"].to_numpy(dtype=float)
    fixed_effects = panel[["firm_id", "industry_year_id"]].to_numpy(dtype=np.int64)
    algorithm = create(fixed_effects, drop_singletons=False)
    transformed = algorithm.residualize(np.column_stack([y, x]))
    y_tilde, x_tilde = transformed[:, 0], transformed[:, 1:]
    fit = sm.OLS(y_tilde, x_tilde, hasconst=False).fit(
        cov_type="cluster",
        cov_kwds={"groups": panel["firm_id"].to_numpy(), "use_correction": True},
        use_t=True,
    )
    names = ["Architecture x linear time", "Architecture x post-2022 break"]
    results = pd.DataFrame(
        {
            "parameter": names,
            "estimate": fit.params,
            "std_error": fit.bse,
            "t_stat": fit.tvalues,
            "p_value": fit.pvalues,
            "ci_low": fit.conf_int()[:, 0],
            "ci_high": fit.conf_int()[:, 1],
            "observations": len(panel),
            "firm_clusters": panel["firm_id"].nunique(),
        }
    )
    observed_break = float(fit.params[1])

    firm = panel.drop_duplicates("firm_id").sort_values("firm_id").reset_index(drop=True)
    rng = np.random.default_rng(SEED)
    exposure_matrix = np.empty((len(firm), B), dtype=np.int8)
    for _, stratum in firm.groupby("permutation_stratum", sort=True):
        indices = stratum.index.to_numpy()
        values = stratum["architecture_count"].to_numpy(dtype=np.int8)
        for draw in range(B):
            exposure_matrix[indices, draw] = rng.permutation(values)
    assignments = pd.DataFrame(
        {
            "firm_id": np.tile(firm["firm_id"].to_numpy(dtype=np.int32), B),
            "draw": np.repeat(np.arange(1, B + 1, dtype=np.int16), len(firm)),
            "permuted_architecture_count": exposure_matrix.T.reshape(-1),
        }
    )
    assignment_path = DATA / "stratified_permutation_assignments.parquet"
    assignments.to_parquet(assignment_path, index=False)

    permutation_trend = np.empty(B, dtype=float)
    permutation_break = np.empty(B, dtype=float)
    firm_index = panel["firm_id"].to_numpy(dtype=int)
    for draw in range(B):
        assigned = exposure_matrix[firm_index, draw].astype(float)
        x_draw = np.column_stack([assigned * time, assigned * post])
        x_draw_tilde = algorithm.residualize(x_draw)
        coefficients = solve_columns(y_tilde, x_draw_tilde)
        permutation_trend[draw] = coefficients[0]
        permutation_break[draw] = coefficients[1]
    randomization_p = float(
        (1 + np.sum(np.abs(permutation_break) >= abs(observed_break))) / (B + 1)
    )
    draws = pd.DataFrame(
        {
            "draw": np.arange(1, B + 1),
            "trend_estimate": permutation_trend,
            "post_break_estimate": permutation_break,
        }
    )
    draws.to_csv(DATA / "detrended_permutation_estimates.csv", index=False)
    draws.to_parquet(DATA / "detrended_permutation_estimates.parquet", index=False)
    results.loc[1, "stratified_randomization_p"] = randomization_p
    export_table(results, "table1_detrended_result", "Detrended Cross-Architecture Result")

    event = pd.read_csv(SOURCE_EXPERIMENT / "data" / "event_study_gate.csv")
    delta = float(fit.params[0])
    event["linear_projection"] = event["relative_to_2021"] * delta
    event["deviation_from_linear_projection"] = (
        event["estimate_per_architecture"] - event["linear_projection"]
    )
    event.to_csv(DATA / "event_linear_projection.csv", index=False)
    export_table(event, "table2_event_linear_projection", "Event Coefficients and Linear Projection")

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.plot(
        event["relative_to_2021"],
        event["estimate_per_architecture"],
        marker="o",
        color="#0072B2",
        label="Event estimate",
    )
    ax.plot(
        event["relative_to_2021"],
        event["linear_projection"],
        linestyle="--",
        color="#D55E00",
        label="Estimated linear differential trend",
    )
    ax.axhline(0, color="black", linewidth=0.8)
    ax.axvline(0.5, color="#555555", linestyle=":", linewidth=1)
    ax.set_xlabel("Year relative to 2021")
    ax.set_ylabel("Difference per frozen architecture")
    ax.legend(frameon=False)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.5)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig1_event_vs_linear_trend.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig1_event_vs_linear_trend.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.5, 3.7))
    ax.hist(permutation_break, bins=35, color="#B8B8B8", edgecolor="white")
    ax.axvline(observed_break, color="#D55E00", linewidth=1.4, label="Observed detrended break")
    ax.axvline(-observed_break, color="#D55E00", linewidth=1, linestyle="--")
    ax.set_xlabel("Permuted detrended post break")
    ax.set_ylabel("Frequency")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(FIGURES / "fig2_detrended_permutation.pdf", bbox_inches="tight")
    fig.savefig(FIGURES / "fig2_detrended_permutation.png", dpi=300, bbox_inches="tight")
    plt.close(fig)

    survives = bool(observed_break < 0 and randomization_p < 0.05)
    diagnostics = {
        "experiment_id": "10_cross_architecture_detrended",
        "gate": "DETREND_AND_RANDOMIZATION",
        "source_sha256": sha256(SOURCE),
        "cloned_sha256": sha256(cloned),
        "linear_trend_estimate": float(fit.params[0]),
        "linear_trend_std_error": float(fit.bse[0]),
        "detrended_post_break": observed_break,
        "detrended_post_break_std_error": float(fit.bse[1]),
        "detrended_post_break_p_value": float(fit.pvalues[1]),
        "stratified_randomization_p": randomization_p,
        "draws": B,
        "seed": SEED,
        "assignment_sha256": sha256(assignment_path),
        "experiment_09_result_survives_detrending": survives,
    }
    (LOGS / "detrended_inference.json").write_text(json.dumps(diagnostics, indent=2), encoding="utf-8")
    status = json.loads(STATUS.read_text(encoding="utf-8"))
    status["phase"] = "DETREND_AND_RANDOMIZATION_RUN"
    status["detrended_result_survives"] = survives
    status["next_gate"] = "FINALIZE_ROBUST" if survives else "ARCHIVE_FAILED_SENSITIVITY"
    STATUS.write_text(json.dumps(status, indent=2), encoding="utf-8")
    print(json.dumps(diagnostics, indent=2))


if __name__ == "__main__":
    main()
